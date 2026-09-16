# app.py
import os
from datetime import datetime, date, timedelta
from functools import wraps

from flask import Flask, render_template, request, redirect, url_for, session, flash, jsonify, send_file
from flask_sqlalchemy import SQLAlchemy
from flask_login import LoginManager, login_user, logout_user, login_required, current_user
from werkzeug.security import generate_password_hash, check_password_hash
from dotenv import load_dotenv
from sqlalchemy import or_, and_

load_dotenv()

from models import (
    db, User, Category, Size, Thickness, Supplier, Customer,
    FactoryRawMaterial, FactoryProduction, FactoryDiary,
    StoreSale, StorePurchase, StoreInventory, StoreReceiving, StoreReturn, StoreDiary,
    StoreSaleItem, StorePurchaseItem, Payment,
    TreasuryAccount, TreasuryTransaction, TreasuryTransfer,
    EditLog, Notification, ActivityLog,
    InventoryAudit,
    AlMasaCrane, AlMasaPartner, AlMasaOperation, AlMasaCheck, AlMasaCheckOperation, AlMasaExpense,
    AlMasaPrivateCrane, AlMasaPrivateOperation, AlMasaPrivateCheck, AlMasaPrivateCheckOperation, AlMasaPrivateExpense,
    AlMasaSupply, AlMasaSupplyOperation, AlMasaSupplyExpense,
    AlMasaPartnerAccount,
    ChatMessage, ChatGroup, ChatGroupMember
)
from utils import (
    login_required as custom_login_required,
    role_required,
    can_edit,
    add_row_to_excel,
    generate_word_report,
    GoogleDriveService
)

app = Flask(__name__)

# ==================== إعدادات الجلسة ====================
app.config['SECRET_KEY'] = os.environ.get('SECRET_KEY', 'dev-secret-key-change-me')
app.config['SESSION_TYPE'] = 'filesystem'
app.config['SESSION_FILE_DIR'] = '/tmp/flask_session'
app.config['SESSION_PERMANENT'] = True
app.config['SESSION_USE_SIGNER'] = True
app.config['SESSION_COOKIE_SAMESITE'] = 'Lax'
app.config['SESSION_COOKIE_SECURE'] = False
app.config['PERMANENT_SESSION_LIFETIME'] = timedelta(days=30)

# ==================== قاعدة البيانات ====================
database_url = os.environ.get('DATABASE_URL', 'sqlite:///instance/factory.db')
if database_url.startswith('postgres://'):
    database_url = database_url.replace('postgres://', 'postgresql://', 1)

app.config['SQLALCHEMY_DATABASE_URI'] = database_url
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False
app.config['SQLALCHEMY_ENGINE_OPTIONS'] = {
    'pool_pre_ping': True,
    'pool_timeout': 30,
    'pool_recycle': 300
}

db.init_app(app)

# ==================== Flask-Login ====================
login_manager = LoginManager()
login_manager.init_app(app)
login_manager.login_view = 'login'
login_manager.login_message = 'يرجى تسجيل الدخول أولاً'

@login_manager.user_loader
def load_user(user_id):
    return User.query.get(int(user_id))

# ==================== خدمة Google Drive ====================
drive_service = GoogleDriveService(
    credentials_file=os.environ.get('GOOGLE_CREDENTIALS_FILE', 'client_secrets.json'),
    folder_id=os.environ.get('GOOGLE_DRIVE_FOLDER_ID', '')
)

# ==================== دوال مساعدة ====================
def get_or_create_treasury_account(person_name, account_type):
    account = TreasuryAccount.query.filter_by(person_name=person_name, account_type=account_type).first()
    if not account:
        account = TreasuryAccount(person_name=person_name, account_type=account_type, balance=0)
        db.session.add(account)
        db.session.commit()
    return account

def get_visible_accounts_for_current_user():
    if current_user.role in ['meg', 'admin', 'mariam', 'sayed']:
        return TreasuryAccount.query.all()
    elif current_user.role == 'ahmed':
        return TreasuryAccount.query.filter_by(person_name='الحاج أحمد').all()
    elif current_user.role == 'eid':
        return TreasuryAccount.query.filter_by(person_name='عيد').all()
    elif current_user.role == 'abdo':
        return TreasuryAccount.query.filter_by(person_name='عبدالله').all()
    else:
        return []

def can_delete_record(user_role, record_date=None):
    if user_role in ['meg', 'admin', 'mariam', 'sayed']:
        return True
    if user_role == 'rehab':
        if record_date:
            return (datetime.now().date() - record_date).days <= 7
        return False
    return False

def log_activity(user_id, action, details=''):
    try:
        activity = ActivityLog(user_id=user_id, action=action, details=details, timestamp=datetime.utcnow())
        db.session.add(activity)
        user = User.query.get(user_id)
        if user:
            user.last_activity = datetime.utcnow()
            db.session.add(user)
        db.session.commit()
    except Exception as e:
        print(f"Error logging activity: {e}")


# ==================== تهيئة قاعدة البيانات ====================
def init_db():
    with app.app_context():
        db.create_all()
        
        # Migration للمستخدمين
        try:
            db.session.execute(db.text('ALTER TABLE users ADD COLUMN IF NOT EXISTS phone VARCHAR(20)'))
            db.session.execute(db.text('ALTER TABLE users ADD COLUMN IF NOT EXISTS last_activity TIMESTAMP'))
            db.session.execute(db.text('ALTER TABLE store_receiving ADD COLUMN IF NOT EXISTS supplier VARCHAR(100)'))
            db.session.commit()
        except Exception as e:
            db.session.rollback()
        
        # ==================== Migration لعمليات الونش (مشاركة) ====================
        try:
            # الحقول القديمة
            db.session.execute(db.text('ALTER TABLE almasa_operations ADD COLUMN IF NOT EXISTS rental_value FLOAT DEFAULT 0'))
            db.session.execute(db.text('ALTER TABLE almasa_operations ADD COLUMN IF NOT EXISTS supply_value FLOAT DEFAULT 0'))
            db.session.execute(db.text('ALTER TABLE almasa_operations ADD COLUMN IF NOT EXISTS rental_total FLOAT DEFAULT 0'))
            db.session.execute(db.text('ALTER TABLE almasa_operations ADD COLUMN IF NOT EXISTS supply_total FLOAT DEFAULT 0'))
            db.session.execute(db.text('ALTER TABLE almasa_operations ADD COLUMN IF NOT EXISTS extra_hours FLOAT DEFAULT 0'))
            db.session.execute(db.text('ALTER TABLE almasa_operations ADD COLUMN IF NOT EXISTS hour_rate FLOAT DEFAULT 0'))
            db.session.execute(db.text('ALTER TABLE almasa_operations ADD COLUMN IF NOT EXISTS travel_days FLOAT DEFAULT 0'))
            db.session.execute(db.text('ALTER TABLE almasa_operations ADD COLUMN IF NOT EXISTS travel_rate FLOAT DEFAULT 0'))
            db.session.execute(db.text('ALTER TABLE almasa_operations ADD COLUMN IF NOT EXISTS travel_supply_rate FLOAT DEFAULT 0'))
            db.session.execute(db.text('ALTER TABLE almasa_operations ADD COLUMN IF NOT EXISTS travel_rental_rate FLOAT DEFAULT 0'))
            db.session.execute(db.text('ALTER TABLE almasa_operations ADD COLUMN IF NOT EXISTS tax_14_enabled BOOLEAN DEFAULT FALSE'))
            db.session.execute(db.text('ALTER TABLE almasa_operations ADD COLUMN IF NOT EXISTS tax_14_value FLOAT DEFAULT 14'))
            db.session.execute(db.text('ALTER TABLE almasa_operations ADD COLUMN IF NOT EXISTS tax_85_enabled BOOLEAN DEFAULT FALSE'))
            db.session.execute(db.text('ALTER TABLE almasa_operations ADD COLUMN IF NOT EXISTS tax_85_value FLOAT DEFAULT 8.5'))
            db.session.execute(db.text('ALTER TABLE almasa_operations ADD COLUMN IF NOT EXISTS invoice_number VARCHAR(50)'))
            db.session.execute(db.text('ALTER TABLE almasa_operations ADD COLUMN IF NOT EXISTS invoice_date DATE'))
            db.session.execute(db.text('ALTER TABLE almasa_operations ADD COLUMN IF NOT EXISTS project_name VARCHAR(200)'))
            db.session.execute(db.text('ALTER TABLE almasa_operations ADD COLUMN IF NOT EXISTS project_location VARCHAR(200)'))
            db.session.execute(db.text('ALTER TABLE almasa_operations ADD COLUMN IF NOT EXISTS payment_method VARCHAR(20)'))
            db.session.execute(db.text('ALTER TABLE almasa_operations ADD COLUMN IF NOT EXISTS check_received BOOLEAN DEFAULT FALSE'))
            db.session.execute(db.text('ALTER TABLE almasa_operations ADD COLUMN IF NOT EXISTS check_received_date DATE'))
            # ✅ الحقول الجديدة
            db.session.execute(db.text('ALTER TABLE almasa_operations ADD COLUMN IF NOT EXISTS rental_days FLOAT DEFAULT 0'))
            db.session.execute(db.text('ALTER TABLE almasa_operations ADD COLUMN IF NOT EXISTS supply_days FLOAT DEFAULT 0'))
            db.session.execute(db.text('ALTER TABLE almasa_operations ADD COLUMN IF NOT EXISTS rental_extra_hours FLOAT DEFAULT 0'))
            db.session.execute(db.text('ALTER TABLE almasa_operations ADD COLUMN IF NOT EXISTS rental_hour_rate FLOAT DEFAULT 0'))
            db.session.execute(db.text('ALTER TABLE almasa_operations ADD COLUMN IF NOT EXISTS supply_extra_hours FLOAT DEFAULT 0'))
            db.session.execute(db.text('ALTER TABLE almasa_operations ADD COLUMN IF NOT EXISTS supply_hour_rate FLOAT DEFAULT 0'))
            db.session.execute(db.text('ALTER TABLE almasa_operations ADD COLUMN IF NOT EXISTS rental_travel_days FLOAT DEFAULT 0'))
            db.session.execute(db.text('ALTER TABLE almasa_operations ADD COLUMN IF NOT EXISTS supply_travel_days FLOAT DEFAULT 0'))
            db.session.execute(db.text('ALTER TABLE almasa_operations ADD COLUMN IF NOT EXISTS payment_check_number VARCHAR(50)'))
            db.session.execute(db.text('ALTER TABLE almasa_operations ADD COLUMN IF NOT EXISTS payment_check_due_date DATE'))
            db.session.execute(db.text('ALTER TABLE almasa_operations ADD COLUMN IF NOT EXISTS payment_account_name VARCHAR(100)'))
            db.session.execute(db.text('ALTER TABLE almasa_operations ADD COLUMN IF NOT EXISTS payment_account_number VARCHAR(50)'))
            db.session.execute(db.text('ALTER TABLE almasa_operations ADD COLUMN IF NOT EXISTS payment_date DATE'))
            db.session.commit()
            print("✅ تم إضافة كل حقول almasa_operations")
        except Exception as e:
            db.session.rollback()
            print(f"⚠️ خطأ almasa_operations: {e}")
        
        # نقل البيانات القديمة
        try:
            db.session.execute(db.text('''
                UPDATE almasa_operations 
                SET rental_value = actual_daily_value,
                    supply_value = default_daily_value,
                    rental_total = actual_total,
                    supply_total = default_total
                WHERE (rental_value IS NULL OR rental_value = 0) 
                  AND actual_daily_value IS NOT NULL
            '''))
            db.session.commit()
            print("✅ تم نقل بيانات almasa_operations")
        except Exception as e:
            db.session.rollback()
            print(f"⚠️ نقل بيانات almasa_operations: {e}")
        
        # ==================== Migration لعمليات الونش الخاص ====================
        try:
            db.session.execute(db.text('ALTER TABLE almasa_private_operations ADD COLUMN IF NOT EXISTS rental_value FLOAT DEFAULT 0'))
            db.session.execute(db.text('ALTER TABLE almasa_private_operations ADD COLUMN IF NOT EXISTS supply_value FLOAT DEFAULT 0'))
            db.session.execute(db.text('ALTER TABLE almasa_private_operations ADD COLUMN IF NOT EXISTS rental_total FLOAT DEFAULT 0'))
            db.session.execute(db.text('ALTER TABLE almasa_private_operations ADD COLUMN IF NOT EXISTS supply_total FLOAT DEFAULT 0'))
            db.session.execute(db.text('ALTER TABLE almasa_private_operations ADD COLUMN IF NOT EXISTS extra_hours FLOAT DEFAULT 0'))
            db.session.execute(db.text('ALTER TABLE almasa_private_operations ADD COLUMN IF NOT EXISTS hour_rate FLOAT DEFAULT 0'))
            db.session.execute(db.text('ALTER TABLE almasa_private_operations ADD COLUMN IF NOT EXISTS travel_days FLOAT DEFAULT 0'))
            db.session.execute(db.text('ALTER TABLE almasa_private_operations ADD COLUMN IF NOT EXISTS travel_rate FLOAT DEFAULT 0'))
            db.session.execute(db.text('ALTER TABLE almasa_private_operations ADD COLUMN IF NOT EXISTS travel_supply_rate FLOAT DEFAULT 0'))
            db.session.execute(db.text('ALTER TABLE almasa_private_operations ADD COLUMN IF NOT EXISTS travel_rental_rate FLOAT DEFAULT 0'))
            db.session.execute(db.text('ALTER TABLE almasa_private_operations ADD COLUMN IF NOT EXISTS tax_14_enabled BOOLEAN DEFAULT FALSE'))
            db.session.execute(db.text('ALTER TABLE almasa_private_operations ADD COLUMN IF NOT EXISTS tax_14_value FLOAT DEFAULT 14'))
            db.session.execute(db.text('ALTER TABLE almasa_private_operations ADD COLUMN IF NOT EXISTS tax_85_enabled BOOLEAN DEFAULT FALSE'))
            db.session.execute(db.text('ALTER TABLE almasa_private_operations ADD COLUMN IF NOT EXISTS tax_85_value FLOAT DEFAULT 8.5'))
            db.session.execute(db.text('ALTER TABLE almasa_private_operations ADD COLUMN IF NOT EXISTS invoice_number VARCHAR(50)'))
            db.session.execute(db.text('ALTER TABLE almasa_private_operations ADD COLUMN IF NOT EXISTS invoice_date DATE'))
            db.session.execute(db.text('ALTER TABLE almasa_private_operations ADD COLUMN IF NOT EXISTS project_name VARCHAR(200)'))
            db.session.execute(db.text('ALTER TABLE almasa_private_operations ADD COLUMN IF NOT EXISTS project_location VARCHAR(200)'))
            db.session.execute(db.text('ALTER TABLE almasa_private_operations ADD COLUMN IF NOT EXISTS payment_method VARCHAR(20)'))
            db.session.execute(db.text('ALTER TABLE almasa_private_operations ADD COLUMN IF NOT EXISTS check_received BOOLEAN DEFAULT FALSE'))
            db.session.execute(db.text('ALTER TABLE almasa_private_operations ADD COLUMN IF NOT EXISTS check_received_date DATE'))
            # ✅ الحقول الجديدة
            db.session.execute(db.text('ALTER TABLE almasa_private_operations ADD COLUMN IF NOT EXISTS rental_days FLOAT DEFAULT 0'))
            db.session.execute(db.text('ALTER TABLE almasa_private_operations ADD COLUMN IF NOT EXISTS supply_days FLOAT DEFAULT 0'))
            db.session.execute(db.text('ALTER TABLE almasa_private_operations ADD COLUMN IF NOT EXISTS rental_extra_hours FLOAT DEFAULT 0'))
            db.session.execute(db.text('ALTER TABLE almasa_private_operations ADD COLUMN IF NOT EXISTS rental_hour_rate FLOAT DEFAULT 0'))
            db.session.execute(db.text('ALTER TABLE almasa_private_operations ADD COLUMN IF NOT EXISTS supply_extra_hours FLOAT DEFAULT 0'))
            db.session.execute(db.text('ALTER TABLE almasa_private_operations ADD COLUMN IF NOT EXISTS supply_hour_rate FLOAT DEFAULT 0'))
            db.session.execute(db.text('ALTER TABLE almasa_private_operations ADD COLUMN IF NOT EXISTS rental_travel_days FLOAT DEFAULT 0'))
            db.session.execute(db.text('ALTER TABLE almasa_private_operations ADD COLUMN IF NOT EXISTS supply_travel_days FLOAT DEFAULT 0'))
            db.session.execute(db.text('ALTER TABLE almasa_private_operations ADD COLUMN IF NOT EXISTS payment_check_number VARCHAR(50)'))
            db.session.execute(db.text('ALTER TABLE almasa_private_operations ADD COLUMN IF NOT EXISTS payment_check_due_date DATE'))
            db.session.execute(db.text('ALTER TABLE almasa_private_operations ADD COLUMN IF NOT EXISTS payment_account_name VARCHAR(100)'))
            db.session.execute(db.text('ALTER TABLE almasa_private_operations ADD COLUMN IF NOT EXISTS payment_account_number VARCHAR(50)'))
            db.session.execute(db.text('ALTER TABLE almasa_private_operations ADD COLUMN IF NOT EXISTS payment_date DATE'))
            db.session.commit()
            print("✅ تم إضافة كل حقول almasa_private_operations")
        except Exception as e:
            db.session.rollback()
            print(f"⚠️ خطأ almasa_private_operations: {e}")
        
        try:
            db.session.execute(db.text('''
                UPDATE almasa_private_operations 
                SET rental_value = actual_daily_value,
                    supply_value = default_daily_value,
                    rental_total = actual_total,
                    supply_total = default_total
                WHERE (rental_value IS NULL OR rental_value = 0) 
                  AND actual_daily_value IS NOT NULL
            '''))
            db.session.commit()
            print("✅ تم نقل بيانات almasa_private_operations")
        except Exception as e:
            db.session.rollback()
            print(f"⚠️ نقل بيانات almasa_private_operations: {e}")
        
        # ==================== Migration لعمليات التوريدات ====================
        try:
            db.session.execute(db.text('ALTER TABLE almasa_supply_operations ADD COLUMN IF NOT EXISTS rental_value FLOAT DEFAULT 0'))
            db.session.execute(db.text('ALTER TABLE almasa_supply_operations ADD COLUMN IF NOT EXISTS supply_value FLOAT DEFAULT 0'))
            db.session.execute(db.text('ALTER TABLE almasa_supply_operations ADD COLUMN IF NOT EXISTS rental_total FLOAT DEFAULT 0'))
            db.session.execute(db.text('ALTER TABLE almasa_supply_operations ADD COLUMN IF NOT EXISTS supply_total FLOAT DEFAULT 0'))
            db.session.execute(db.text('ALTER TABLE almasa_supply_operations ADD COLUMN IF NOT EXISTS extra_hours FLOAT DEFAULT 0'))
            db.session.execute(db.text('ALTER TABLE almasa_supply_operations ADD COLUMN IF NOT EXISTS hour_rate FLOAT DEFAULT 0'))
            db.session.execute(db.text('ALTER TABLE almasa_supply_operations ADD COLUMN IF NOT EXISTS travel_days FLOAT DEFAULT 0'))
            db.session.execute(db.text('ALTER TABLE almasa_supply_operations ADD COLUMN IF NOT EXISTS travel_rate FLOAT DEFAULT 0'))
            db.session.execute(db.text('ALTER TABLE almasa_supply_operations ADD COLUMN IF NOT EXISTS travel_supply_rate FLOAT DEFAULT 0'))
            db.session.execute(db.text('ALTER TABLE almasa_supply_operations ADD COLUMN IF NOT EXISTS travel_rental_rate FLOAT DEFAULT 0'))
            db.session.execute(db.text('ALTER TABLE almasa_supply_operations ADD COLUMN IF NOT EXISTS tax_14_enabled BOOLEAN DEFAULT FALSE'))
            db.session.execute(db.text('ALTER TABLE almasa_supply_operations ADD COLUMN IF NOT EXISTS tax_14_value FLOAT DEFAULT 14'))
            db.session.execute(db.text('ALTER TABLE almasa_supply_operations ADD COLUMN IF NOT EXISTS tax_85_enabled BOOLEAN DEFAULT FALSE'))
            db.session.execute(db.text('ALTER TABLE almasa_supply_operations ADD COLUMN IF NOT EXISTS tax_85_value FLOAT DEFAULT 8.5'))
            db.session.execute(db.text('ALTER TABLE almasa_supply_operations ADD COLUMN IF NOT EXISTS invoice_number VARCHAR(50)'))
            db.session.execute(db.text('ALTER TABLE almasa_supply_operations ADD COLUMN IF NOT EXISTS invoice_date DATE'))
            db.session.execute(db.text('ALTER TABLE almasa_supply_operations ADD COLUMN IF NOT EXISTS project_name VARCHAR(200)'))
            db.session.execute(db.text('ALTER TABLE almasa_supply_operations ADD COLUMN IF NOT EXISTS project_location VARCHAR(200)'))
            db.session.execute(db.text('ALTER TABLE almasa_supply_operations ADD COLUMN IF NOT EXISTS payment_method VARCHAR(20)'))
            db.session.execute(db.text('ALTER TABLE almasa_supply_operations ADD COLUMN IF NOT EXISTS check_received BOOLEAN DEFAULT FALSE'))
            db.session.execute(db.text('ALTER TABLE almasa_supply_operations ADD COLUMN IF NOT EXISTS check_received_date DATE'))
            # ✅ الحقول الجديدة
            db.session.execute(db.text('ALTER TABLE almasa_supply_operations ADD COLUMN IF NOT EXISTS rental_days FLOAT DEFAULT 0'))
            db.session.execute(db.text('ALTER TABLE almasa_supply_operations ADD COLUMN IF NOT EXISTS supply_days FLOAT DEFAULT 0'))
            db.session.execute(db.text('ALTER TABLE almasa_supply_operations ADD COLUMN IF NOT EXISTS rental_extra_hours FLOAT DEFAULT 0'))
            db.session.execute(db.text('ALTER TABLE almasa_supply_operations ADD COLUMN IF NOT EXISTS rental_hour_rate FLOAT DEFAULT 0'))
            db.session.execute(db.text('ALTER TABLE almasa_supply_operations ADD COLUMN IF NOT EXISTS supply_extra_hours FLOAT DEFAULT 0'))
            db.session.execute(db.text('ALTER TABLE almasa_supply_operations ADD COLUMN IF NOT EXISTS supply_hour_rate FLOAT DEFAULT 0'))
            db.session.execute(db.text('ALTER TABLE almasa_supply_operations ADD COLUMN IF NOT EXISTS rental_travel_days FLOAT DEFAULT 0'))
            db.session.execute(db.text('ALTER TABLE almasa_supply_operations ADD COLUMN IF NOT EXISTS supply_travel_days FLOAT DEFAULT 0'))
            db.session.execute(db.text('ALTER TABLE almasa_supply_operations ADD COLUMN IF NOT EXISTS payment_check_number VARCHAR(50)'))
            db.session.execute(db.text('ALTER TABLE almasa_supply_operations ADD COLUMN IF NOT EXISTS payment_check_due_date DATE'))
            db.session.execute(db.text('ALTER TABLE almasa_supply_operations ADD COLUMN IF NOT EXISTS payment_account_name VARCHAR(100)'))
            db.session.execute(db.text('ALTER TABLE almasa_supply_operations ADD COLUMN IF NOT EXISTS payment_account_number VARCHAR(50)'))
            db.session.execute(db.text('ALTER TABLE almasa_supply_operations ADD COLUMN IF NOT EXISTS payment_date DATE'))
            db.session.commit()
            print("✅ تم إضافة كل حقول almasa_supply_operations")
        except Exception as e:
            db.session.rollback()
            print(f"⚠️ خطأ almasa_supply_operations: {e}")
        
        try:
            db.session.execute(db.text('''
                UPDATE almasa_supply_operations 
                SET rental_value = customer_daily_rate,
                    supply_value = supplier_daily_rate,
                    rental_total = customer_total,
                    supply_total = supplier_total
                WHERE (rental_value IS NULL OR rental_value = 0) 
                  AND customer_daily_rate IS NOT NULL
            '''))
            db.session.commit()
            print("✅ تم نقل بيانات almasa_supply_operations")
        except Exception as e:
            db.session.rollback()
            print(f"⚠️ نقل بيانات almasa_supply_operations: {e}")
        
        # ==================== Migration للمصاريف ====================
        try:
            db.session.execute(db.text('ALTER TABLE almasa_expenses ADD COLUMN IF NOT EXISTS operation_id INTEGER'))
            db.session.execute(db.text('ALTER TABLE almasa_expenses ADD COLUMN IF NOT EXISTS paid_by VARCHAR(100)'))
            db.session.commit()
        except Exception as e:
            db.session.rollback()
        
        try:
            db.session.execute(db.text('ALTER TABLE almasa_private_expenses ADD COLUMN IF NOT EXISTS operation_id INTEGER'))
            db.session.execute(db.text('ALTER TABLE almasa_private_expenses ADD COLUMN IF NOT EXISTS paid_by VARCHAR(100)'))
            db.session.commit()
        except Exception as e:
            db.session.rollback()
        
        try:
            db.session.execute(db.text('ALTER TABLE almasa_supply_expenses ADD COLUMN IF NOT EXISTS operation_id INTEGER'))
            db.session.execute(db.text('ALTER TABLE almasa_supply_expenses ADD COLUMN IF NOT EXISTS paid_by VARCHAR(100)'))
            db.session.commit()
        except Exception as e:
            db.session.rollback()
        
        # ==================== Migration للشيكات ====================
        try:
            db.session.execute(db.text('ALTER TABLE almasa_checks ADD COLUMN IF NOT EXISTS operation_id INTEGER'))
            db.session.execute(db.text('ALTER TABLE almasa_checks ADD COLUMN IF NOT EXISTS issue_date DATE'))
            db.session.commit()
        except Exception as e:
            db.session.rollback()
        
        try:
            db.session.execute(db.text('ALTER TABLE almasa_private_checks ADD COLUMN IF NOT EXISTS operation_id INTEGER'))
            db.session.execute(db.text('ALTER TABLE almasa_private_checks ADD COLUMN IF NOT EXISTS issue_date DATE'))
            db.session.commit()
        except Exception as e:
            db.session.rollback()
        
        # ==================== Migration للشات ====================
        try:
            db.session.execute(db.text('ALTER TABLE chat_messages ADD COLUMN IF NOT EXISTS is_deleted BOOLEAN DEFAULT FALSE'))
            db.session.execute(db.text('ALTER TABLE chat_messages ADD COLUMN IF NOT EXISTS is_read BOOLEAN DEFAULT FALSE'))
            db.session.execute(db.text('ALTER TABLE chat_messages ADD COLUMN IF NOT EXISTS group_id INTEGER'))
            db.session.execute(db.text('ALTER TABLE chat_messages ADD COLUMN IF NOT EXISTS recipient_id INTEGER'))
            db.session.commit()
            print("✅ تم إضافة حقول chat_messages")
        except Exception as e:
            db.session.rollback()
        
        # ==================== Migration لـ days_count → Float ====================
        try:
            db.session.execute(db.text('ALTER TABLE almasa_operations ALTER COLUMN days_count TYPE FLOAT'))
            db.session.commit()
            print("✅ days_count → Float (operations)")
        except Exception as e:
            db.session.rollback()
            print(f"⚠️ operations days_count: {e}")

        try:
            db.session.execute(db.text('ALTER TABLE almasa_private_operations ALTER COLUMN days_count TYPE FLOAT'))
            db.session.commit()
            print("✅ days_count → Float (private)")
        except Exception as e:
            db.session.rollback()
            print(f"⚠️ private days_count: {e}")

        try:
            db.session.execute(db.text('ALTER TABLE almasa_supply_operations ALTER COLUMN days_count TYPE FLOAT'))
            db.session.commit()
            print("✅ days_count → Float (supply)")
        except Exception as e:
            db.session.rollback()
            print(f"⚠️ supply days_count: {e}")
        
        # ==================== Migration للبيانات القديمة (للحقول الجديدة) ====================
        try:
            db.session.execute(db.text('''
                UPDATE almasa_operations 
                SET rental_days = COALESCE(days_count, 0),
                    supply_days = COALESCE(days_count, 0),
                    rental_travel_days = COALESCE(travel_days, 0),
                    supply_travel_days = COALESCE(travel_days, 0),
                    rental_extra_hours = COALESCE(extra_hours, 0),
                    supply_extra_hours = COALESCE(extra_hours, 0),
                    rental_hour_rate = COALESCE(hour_rate, 0),
                    supply_hour_rate = COALESCE(hour_rate, 0)
                WHERE rental_days IS NULL OR rental_days = 0
            '''))
            db.session.commit()
            print("✅ تم نقل البيانات القديمة للحقول الجديدة (operations)")
        except Exception as e:
            db.session.rollback()
            print(f"⚠️ نقل البيانات الجديدة (operations): {e}")

        try:
            db.session.execute(db.text('''
                UPDATE almasa_private_operations 
                SET rental_days = COALESCE(days_count, 0),
                    supply_days = COALESCE(days_count, 0),
                    rental_travel_days = COALESCE(travel_days, 0),
                    supply_travel_days = COALESCE(travel_days, 0),
                    rental_extra_hours = COALESCE(extra_hours, 0),
                    supply_extra_hours = COALESCE(extra_hours, 0),
                    rental_hour_rate = COALESCE(hour_rate, 0),
                    supply_hour_rate = COALESCE(hour_rate, 0)
                WHERE rental_days IS NULL OR rental_days = 0
            '''))
            db.session.commit()
            print("✅ تم نقل البيانات القديمة للحقول الجديدة (private)")
        except Exception as e:
            db.session.rollback()
            print(f"⚠️ نقل البيانات الجديدة (private): {e}")

        try:
            db.session.execute(db.text('''
                UPDATE almasa_supply_operations 
                SET rental_days = COALESCE(days_count, 0),
                    supply_days = COALESCE(days_count, 0),
                    rental_travel_days = COALESCE(travel_days, 0),
                    supply_travel_days = COALESCE(travel_days, 0),
                    rental_extra_hours = COALESCE(extra_hours, 0),
                    supply_extra_hours = COALESCE(extra_hours, 0),
                    rental_hour_rate = COALESCE(hour_rate, 0),
                    supply_hour_rate = COALESCE(hour_rate, 0)
                WHERE rental_days IS NULL OR rental_days = 0
            '''))
            db.session.commit()
            print("✅ تم نقل البيانات القديمة للحقول الجديدة (supply)")
        except Exception as e:
            db.session.rollback()
            print(f"⚠️ نقل البيانات الجديدة (supply): {e}")
        
        # ==================== إنشاء المستخدمين ====================
        users_data = [
            {'username': 'meg', 'password': '262004', 'full_name': 'MEG', 'role': 'meg', 'is_hidden': True},
            {'username': 'f', 'password': '*1997#', 'full_name': 'Admin', 'role': 'admin', 'is_hidden': False},
            {'username': 'mariam', 'password': '#mariam2004', 'full_name': 'Mariam', 'role': 'mariam', 'is_hidden': False},
            {'username': 'rehab', 'password': 'rehab2004#', 'full_name': 'Rehab', 'role': 'rehab', 'is_hidden': False},
            {'username': 'mohamed', 'password': 'mohamed123#', 'full_name': 'Mohamed', 'role': 'mohamed', 'is_hidden': False},
            {'username': 'a', 'password': '#123456#', 'full_name': 'الحاج أحمد', 'role': 'ahmed', 'is_hidden': False},
            {'username': 'eid', 'password': 'eid123#', 'full_name': 'عيد', 'role': 'eid', 'is_hidden': False},
            {'username': 'abdo', 'password': 'abdo123#', 'full_name': 'عبدالله', 'role': 'abdo', 'is_hidden': False},
            {'username': 'sayed', 'password': 'sayed1977#', 'full_name': 'سيد', 'role': 'sayed', 'is_hidden': False},
            {'username': 'dina', 'password': 'dina2003', 'full_name': 'دينا', 'role': 'dina', 'is_hidden': False},
        ]

        for user_data in users_data:
            existing_user = User.query.filter_by(username=user_data['username']).first()
            if not existing_user:
                new_user = User(
                    username=user_data['username'],
                    full_name=user_data['full_name'],
                    role=user_data['role'],
                    is_hidden=user_data['is_hidden'],
                    created_at=datetime.utcnow()
                )
                new_user.set_password(user_data['password'])
                db.session.add(new_user)

        db.session.commit()
        print("✅ تم تهيئة قاعدة البيانات وإنشاء المستخدمين")

        treasury_persons = ['الحاج أحمد', 'عيد', 'عبدالله', 'الحاج فتحي']
        account_types = ['كاش', 'فودافون كاش', 'انستا باي', 'شيك']
        for person in treasury_persons:
            for acc_type in account_types:
                get_or_create_treasury_account(person, acc_type)

        print("✅ تم إنشاء حسابات الخزينة")
        
init_db()
# ==================== Context Processor ====================
@app.context_processor
def inject_globals():
    unread_notifications = 0
    unread_messages_count = 0
    if current_user.is_authenticated:
        unread_notifications = Notification.query.filter_by(user_id=current_user.id, is_read=False).count()
        unread_messages_count = ChatMessage.query.filter(
            ChatMessage.recipient_id == current_user.id,
            ChatMessage.is_read == False,
            ChatMessage.is_deleted == False
        ).count()
    def get_user_name(user_id):
        if not user_id:
            return 'غير معروف'
        user = User.query.get(user_id)
        return user.full_name if user else 'غير معروف'
    return {
        'now': datetime.now(),
        'unread_notifications': unread_notifications,
        'unread_messages_count': unread_messages_count,
        'get_user_name': get_user_name
    }

# ==================== الصفحات الأساسية ====================
@app.route('/health')
def health():
    return jsonify({'status': 'ok', 'message': 'مصنع الفتح شغال'})

@app.route('/')
def index():
    if current_user.is_authenticated:
        return redirect(url_for('dashboard'))
    return redirect(url_for('login'))

@app.route('/login', methods=['GET', 'POST'])
def login():
    if current_user.is_authenticated:
        return redirect(url_for('dashboard'))
    if request.method == 'POST':
        username = request.form.get('username')
        password = request.form.get('password')
        user = User.query.filter_by(username=username).first()
        if user and user.check_password(password):
            login_user(user)
            session.permanent = True
            session['role'] = user.role
            session['user_id'] = user.id
            session['full_name'] = user.full_name
            log_activity(user.id, 'login', f"تسجيل دخول {user.full_name}")
            flash(f'مرحباً {user.full_name} 👋', 'success')
            return redirect(url_for('dashboard'))
        else:
            flash('اسم المستخدم أو كلمة المرور غير صحيحة', 'danger')
    return render_template('login.html')

@app.route('/choose-company')
@custom_login_required
def choose_company():
    if current_user.role != 'sayed':
        return redirect(url_for('dashboard'))
    return render_template('choose_company.html')

@app.route('/logout')
@login_required
def logout():
    log_activity(current_user.id, 'logout', f"تسجيل خروج {current_user.full_name}")
    logout_user()
    session.clear()
    flash('تم تسجيل الخروج بنجاح', 'info')
    return redirect(url_for('login'))

@app.route('/dashboard')
@custom_login_required
def dashboard():
    role = current_user.role
    if role == 'dina':
        return redirect(url_for('almasa_index'))
    
    stats = {
        'raw_materials_count': FactoryRawMaterial.query.count() if role in ['meg','admin','mariam','rehab','mohamed','sayed'] else 0,
        'production_count': FactoryProduction.query.count() if role in ['meg','admin','mariam','rehab','mohamed','sayed'] else 0,
        'sales_count': StoreSale.query.count() if role in ['meg','admin','mariam','rehab','ahmed','sayed'] else 0,
        'purchases_count': StorePurchase.query.count() if role in ['meg','admin','mariam','rehab','ahmed','sayed'] else 0,
        'customers_count': Customer.query.count() if role in ['meg','admin','mariam','rehab','ahmed','sayed'] else 0,
        'suppliers_count': Supplier.query.count() if role in ['meg','admin','mariam','rehab','ahmed','sayed'] else 0,
        'treasury_balance': 0,
        'today_sales': db.session.query(db.func.sum(StoreSaleItem.total)).join(StoreSale).filter(StoreSale.date == date.today()).scalar() or 0,
        'today_purchases': db.session.query(db.func.sum(StorePurchaseItem.total)).join(StorePurchase).filter(StorePurchase.date == date.today()).scalar() or 0,
        'low_inventory_count': StoreInventory.query.filter(StoreInventory.current_quantity <= StoreInventory.min_quantity).count() if role in ['meg','admin','mariam','rehab','ahmed','sayed'] else 0,
    }

    if role in ['meg', 'admin', 'mariam', 'sayed']:
        stats['treasury_balance'] = db.session.query(db.func.sum(TreasuryAccount.balance)).scalar() or 0
        recent_sales = StoreSale.query.order_by(StoreSale.date.desc()).limit(5).all()
        recent_production = FactoryProduction.query.order_by(FactoryProduction.date.desc()).limit(5).all()
        recent_transactions = TreasuryTransaction.query.order_by(TreasuryTransaction.date.desc()).limit(5).all()
    elif role == 'rehab':
        stats['treasury_balance'] = 0
        recent_sales = StoreSale.query.order_by(StoreSale.date.desc()).limit(5).all()
        recent_production = FactoryProduction.query.order_by(FactoryProduction.date.desc()).limit(5).all()
        recent_transactions = []
    elif role == 'mohamed':
        stats['treasury_balance'] = 0
        recent_sales = []
        recent_production = FactoryProduction.query.order_by(FactoryProduction.date.desc()).limit(5).all()
        recent_transactions = []
    elif role == 'ahmed':
        stats['treasury_balance'] = db.session.query(db.func.sum(TreasuryAccount.balance)).filter(TreasuryAccount.person_name == current_user.full_name).scalar() or 0
        recent_sales = StoreSale.query.order_by(StoreSale.date.desc()).limit(5).all()
        recent_production = []
        recent_transactions = TreasuryTransaction.query.filter_by(created_by=current_user.id).order_by(TreasuryTransaction.date.desc()).limit(5).all()
    else:
        stats['treasury_balance'] = db.session.query(db.func.sum(TreasuryAccount.balance)).filter(TreasuryAccount.person_name == current_user.full_name).scalar() or 0
        recent_sales = []
        recent_production = []
        recent_transactions = TreasuryTransaction.query.filter_by(created_by=current_user.id).order_by(TreasuryTransaction.date.desc()).limit(5).all()

    inactive_users_count = 0
    if role == 'admin':
        threshold = datetime.utcnow() - timedelta(days=3)
        inactive_users_count = User.query.filter(User.is_hidden == False, User.last_activity < threshold).count()

    return render_template('dashboard.html', stats=stats,
                           recent_sales=recent_sales,
                           recent_production=recent_production,
                           recent_transactions=recent_transactions,
                           inactive_users_count=inactive_users_count)

# ==================== المصنع ====================
@app.route('/factory')
@custom_login_required
@role_required('meg', 'admin', 'mariam', 'rehab', 'mohamed', 'sayed')
def factory_index():
    today = date.today()
    raw_materials = FactoryRawMaterial.query.filter_by(date=today).order_by(FactoryRawMaterial.id.asc()).all()
    production = FactoryProduction.query.filter_by(date=today).order_by(FactoryProduction.id.asc()).all()
    diary = FactoryDiary.query.filter_by(date=today).order_by(FactoryDiary.id.asc()).all()
    return render_template('factory/index.html', raw_materials=raw_materials, production=production, diary=diary)

@app.route('/factory/raw-materials', methods=['GET', 'POST'])
@custom_login_required
@role_required('meg', 'admin', 'mariam', 'rehab', 'mohamed', 'sayed')
def factory_raw_materials():
    if request.method == 'POST':
        if request.form.get('delete_id'):
            record_id = int(request.form.get('delete_id'))
            record = FactoryRawMaterial.query.get_or_404(record_id)
            if can_delete_record(current_user.role, record.date):
                db.session.delete(record)
                db.session.commit()
                flash('تم حذف السجل بنجاح', 'success')
            else:
                flash('غير مصرح لك بالحذف', 'danger')
            return redirect(url_for('factory_raw_materials'))
        if request.form.get('edit_id'):
            record_id = int(request.form.get('edit_id'))
            record = FactoryRawMaterial.query.get_or_404(record_id)
            if can_edit(current_user.role, record.date, record.created_by, current_user.id):
                record.date = datetime.strptime(request.form.get('date'), '%Y-%m-%d').date()
                record.pipe_size = request.form.get('pipe_size')
                record.pipe_thickness = request.form.get('pipe_thickness')
                record.quantity = float(request.form.get('quantity', 0))
                record.supplier = request.form.get('supplier')
                record.notes = request.form.get('notes')
                db.session.commit()
                flash('تم تحديث السجل بنجاح', 'success')
            return redirect(url_for('factory_raw_materials'))
        record_date = datetime.strptime(request.form.get('date'), '%Y-%m-%d').date()
        pipe_size = request.form.get('pipe_size')
        pipe_thickness = request.form.get('pipe_thickness')
        quantity = float(request.form.get('quantity', 0))
        supplier = request.form.get('supplier')
        notes = request.form.get('notes')
        if supplier and not Supplier.query.filter_by(name=supplier).first():
            db.session.add(Supplier(name=supplier))
        new_record = FactoryRawMaterial(date=record_date, pipe_size=pipe_size, pipe_thickness=pipe_thickness,
                                        quantity=quantity, supplier=supplier, notes=notes,
                                        created_by=current_user.id, created_at=datetime.utcnow())
        db.session.add(new_record)
        db.session.commit()
        flash('تم إضافة وارد المواسير بنجاح', 'success')
        return redirect(url_for('factory_raw_materials'))
    materials = FactoryRawMaterial.query.order_by(FactoryRawMaterial.date.asc(), FactoryRawMaterial.id.asc()).all()
    suppliers = Supplier.query.order_by(Supplier.name.asc()).all()
    return render_template('factory/raw_materials.html', materials=materials, suppliers=suppliers)

@app.route('/factory/production', methods=['GET', 'POST'])
@custom_login_required
@role_required('meg', 'admin', 'mariam', 'rehab', 'mohamed', 'sayed')
def factory_production():
    if request.method == 'POST':
        if request.form.get('delete_id'):
            record_id = int(request.form.get('delete_id'))
            record = FactoryProduction.query.get_or_404(record_id)
            if can_delete_record(current_user.role, record.date):
                db.session.delete(record)
                db.session.commit()
                flash('تم حذف السجل بنجاح', 'success')
            return redirect(url_for('factory_production'))
        if request.form.get('edit_id'):
            record_id = int(request.form.get('edit_id'))
            record = FactoryProduction.query.get_or_404(record_id)
            if can_edit(current_user.role, record.date, record.created_by, current_user.id):
                record.date = datetime.strptime(request.form.get('date'), '%Y-%m-%d').date()
                record.elbow_size = request.form.get('elbow_size')
                record.elbow_thickness = request.form.get('elbow_thickness')
                record.quantity = float(request.form.get('quantity', 0))
                record.raw_material_used = float(request.form.get('raw_material_used', 0))
                record.notes = request.form.get('notes')
                db.session.commit()
                flash('تم تحديث السجل بنجاح', 'success')
            return redirect(url_for('factory_production'))
        record_date = datetime.strptime(request.form.get('date'), '%Y-%m-%d').date()
        elbow_size = request.form.get('elbow_size')
        elbow_thickness = request.form.get('elbow_thickness')
        quantity = float(request.form.get('quantity', 0))
        raw_material_used = float(request.form.get('raw_material_used', 0))
        notes = request.form.get('notes')
        new_record = FactoryProduction(date=record_date, elbow_size=elbow_size, elbow_thickness=elbow_thickness,
                                       quantity=quantity, raw_material_used=raw_material_used, notes=notes,
                                       created_by=current_user.id, created_at=datetime.utcnow())
        db.session.add(new_record)
        db.session.commit()
        flash('تم تسجيل الإنتاج بنجاح', 'success')
        return redirect(url_for('factory_production'))
    production = FactoryProduction.query.order_by(FactoryProduction.date.asc(), FactoryProduction.id.asc()).all()
    return render_template('factory/production.html', production=production)

@app.route('/factory/diary', methods=['GET', 'POST'])
@custom_login_required
@role_required('meg', 'admin', 'mariam', 'rehab', 'mohamed', 'sayed')
def factory_diary():
    if request.method == 'POST':
        if request.form.get('delete_id'):
            record_id = int(request.form.get('delete_id'))
            record = FactoryDiary.query.get_or_404(record_id)
            if can_delete_record(current_user.role, record.date):
                db.session.delete(record)
                db.session.commit()
                flash('تم حذف السجل بنجاح', 'success')
            return redirect(url_for('factory_diary'))
        if request.form.get('edit_id'):
            record_id = int(request.form.get('edit_id'))
            record = FactoryDiary.query.get_or_404(record_id)
            if can_edit(current_user.role, record.date, record.created_by, current_user.id):
                record.date = datetime.strptime(request.form.get('date'), '%Y-%m-%d').date()
                record.description = request.form.get('description')
                record.amount = float(request.form.get('amount', 0))
                db.session.commit()
                flash('تم تحديث السجل بنجاح', 'success')
            return redirect(url_for('factory_diary'))
        record_date = datetime.strptime(request.form.get('date'), '%Y-%m-%d').date()
        descriptions = request.form.getlist('description[]')
        amounts = request.form.getlist('amount[]')
        for i in range(len(descriptions)):
            if descriptions[i].strip():
                new_record = FactoryDiary(date=record_date, description=descriptions[i],
                                          amount=float(amounts[i]) if i < len(amounts) and amounts[i] else 0,
                                          created_by=current_user.id, created_at=datetime.utcnow())
                db.session.add(new_record)
        db.session.commit()
        flash('تم تسجيل اليومية بنجاح', 'success')
        return redirect(url_for('factory_diary'))
    diary = FactoryDiary.query.order_by(FactoryDiary.date.asc(), FactoryDiary.id.asc()).all()
    return render_template('factory/diary.html', diary=diary)

# ==================== المحل ====================
@app.route('/store')
@custom_login_required
@role_required('meg', 'admin', 'mariam', 'rehab', 'ahmed', 'sayed')
def store_index():
    today = date.today()
    sales = StoreSale.query.filter_by(date=today).order_by(StoreSale.id.asc()).all()
    purchases = StorePurchase.query.filter_by(date=today).order_by(StorePurchase.id.asc()).all()
    receiving = StoreReceiving.query.filter_by(date=today).order_by(StoreReceiving.id.asc()).all()
    inventory_low = StoreInventory.query.filter(StoreInventory.current_quantity <= StoreInventory.min_quantity).count()
    return render_template('store/index.html', sales=sales, purchases=purchases, receiving=receiving, inventory_low=inventory_low)

@app.route('/store/transactions', methods=['GET', 'POST'])
@custom_login_required
@role_required('meg', 'admin', 'mariam', 'rehab', 'ahmed', 'sayed')
def store_transactions():
    if request.method == 'POST':
        if request.form.get('delete_id'):
            record_id = int(request.form.get('delete_id'))
            transaction_type = request.form.get('transaction_type')
            if can_delete_record(current_user.role, StoreSale.query.get_or_404(record_id).date if transaction_type == 'sale' else StorePurchase.query.get_or_404(record_id).date):
                if transaction_type == 'sale':
                    record = StoreSale.query.get_or_404(record_id)
                    for item in record.items:
                        inv = StoreInventory.query.filter_by(product_type=item.product_type,
                                                             product_size=item.product_size,
                                                             product_spec=item.product_spec).first()
                        if inv:
                            inv.current_quantity += item.quantity
                    db.session.delete(record)
                    db.session.commit()
                    flash('تم الحذف بنجاح', 'success')
                elif transaction_type == 'purchase':
                    record = StorePurchase.query.get_or_404(record_id)
                    for item in record.items:
                        inv = StoreInventory.query.filter_by(product_type=item.product_type,
                                                             product_size=item.product_size,
                                                             product_spec=item.product_spec).first()
                        if inv:
                            inv.current_quantity -= item.quantity
                    db.session.delete(record)
                    db.session.commit()
                    flash('تم الحذف بنجاح', 'success')
            return redirect(url_for('store_transactions'))

        if request.form.get('edit_id'):
            record_id = int(request.form.get('edit_id'))
            transaction_type = request.form.get('transaction_type')
            if transaction_type == 'sale':
                record = StoreSale.query.get_or_404(record_id)
                if can_edit(current_user.role, record.date, record.created_by, current_user.id):
                    record.date = datetime.strptime(request.form.get('date'), '%Y-%m-%d').date()
                    record.customer_name = request.form.get('party_name')
                    record.customer_phone = request.form.get('party_phone')
                    record.payment_type = request.form.get('payment_type', 'آجل')
                    for item in record.items:
                        inv = StoreInventory.query.filter_by(product_type=item.product_type,
                                                             product_size=item.product_size,
                                                             product_spec=item.product_spec).first()
                        if inv:
                            inv.current_quantity += item.quantity
                        db.session.delete(item)
                    product_types = request.form.getlist('product_type')
                    product_sizes = request.form.getlist('product_size')
                    product_specs = request.form.getlist('product_spec')
                    quantities = request.form.getlist('quantity')
                    unit_prices = request.form.getlist('unit_price')
                    new_product_types = request.form.getlist('new_product_type')
                    new_product_sizes = request.form.getlist('new_product_size')
                    new_product_specs = request.form.getlist('new_product_spec')
                    for i in range(len(product_types)):
                        if not product_types[i].strip() or not quantities[i].strip():
                            continue
                        product_type = product_types[i]
                        if product_type == 'new':
                            product_type = new_product_types[i] if i < len(new_product_types) else ''
                            if product_type and not Category.query.filter_by(name=product_type).first():
                                db.session.add(Category(name=product_type))
                        product_size = product_sizes[i] if i < len(product_sizes) else ''
                        if product_size == 'new':
                            product_size = new_product_sizes[i] if i < len(new_product_sizes) else ''
                            if product_size and not Size.query.filter_by(value=product_size).first():
                                db.session.add(Size(value=product_size))
                        product_spec = product_specs[i] if i < len(product_specs) else ''
                        if product_spec == 'new':
                            product_spec = new_product_specs[i] if i < len(new_product_specs) else ''
                            if product_spec and not Thickness.query.filter_by(value=product_spec).first():
                                db.session.add(Thickness(value=product_spec))
                        item = StoreSaleItem(
                            sale_id=record.id,
                            product_type=product_type,
                            product_size=product_size,
                            product_spec=product_spec,
                            quantity=float(quantities[i]) if quantities[i] else 0,
                            unit_price=float(unit_prices[i]) if unit_prices[i] else 0,
                        )
                        item.total = item.quantity * item.unit_price
                        db.session.add(item)
                        inv = StoreInventory.query.filter_by(product_type=item.product_type,
                                                             product_size=item.product_size,
                                                             product_spec=item.product_spec).first()
                        if inv:
                            inv.current_quantity -= item.quantity
                        else:
                            inv = StoreInventory(product_type=item.product_type,
                                                 product_size=item.product_size,
                                                 product_spec=item.product_spec,
                                                 current_quantity=-item.quantity)
                            db.session.add(inv)
                    db.session.commit()
                    flash('تم تعديل البيع بنجاح', 'success')
            else:
                record = StorePurchase.query.get_or_404(record_id)
                if can_edit(current_user.role, record.date, record.created_by, current_user.id):
                    record.date = datetime.strptime(request.form.get('date'), '%Y-%m-%d').date()
                    record.supplier_name = request.form.get('party_name')
                    record.supplier_phone = request.form.get('party_phone')
                    record.payment_type = request.form.get('payment_type', 'آجل')
                    for item in record.items:
                        inv = StoreInventory.query.filter_by(product_type=item.product_type,
                                                             product_size=item.product_size,
                                                             product_spec=item.product_spec).first()
                        if inv:
                            inv.current_quantity -= item.quantity
                        db.session.delete(item)
                    product_types = request.form.getlist('product_type')
                    product_sizes = request.form.getlist('product_size')
                    product_specs = request.form.getlist('product_spec')
                    quantities = request.form.getlist('quantity')
                    unit_prices = request.form.getlist('unit_price')
                    new_product_types = request.form.getlist('new_product_type')
                    new_product_sizes = request.form.getlist('new_product_size')
                    new_product_specs = request.form.getlist('new_product_spec')
                    for i in range(len(product_types)):
                        if not product_types[i].strip() or not quantities[i].strip():
                            continue
                        product_type = product_types[i]
                        if product_type == 'new':
                            product_type = new_product_types[i] if i < len(new_product_types) else ''
                            if product_type and not Category.query.filter_by(name=product_type).first():
                                db.session.add(Category(name=product_type))
                        product_size = product_sizes[i] if i < len(product_sizes) else ''
                        if product_size == 'new':
                            product_size = new_product_sizes[i] if i < len(new_product_sizes) else ''
                            if product_size and not Size.query.filter_by(value=product_size).first():
                                db.session.add(Size(value=product_size))
                        product_spec = product_specs[i] if i < len(product_specs) else ''
                        if product_spec == 'new':
                            product_spec = new_product_specs[i] if i < len(new_product_specs) else ''
                            if product_spec and not Thickness.query.filter_by(value=product_spec).first():
                                db.session.add(Thickness(value=product_spec))
                        item = StorePurchaseItem(
                            purchase_id=record.id,
                            product_type=product_type,
                            product_size=product_size,
                            product_spec=product_spec,
                            quantity=float(quantities[i]) if quantities[i] else 0,
                            unit_price=float(unit_prices[i]) if unit_prices[i] else 0,
                        )
                        item.total = item.quantity * item.unit_price
                        db.session.add(item)
                        inv = StoreInventory.query.filter_by(product_type=item.product_type,
                                                             product_size=item.product_size,
                                                             product_spec=item.product_spec).first()
                        if inv:
                            inv.current_quantity += item.quantity
                        else:
                            inv = StoreInventory(product_type=item.product_type,
                                                 product_size=item.product_size,
                                                 product_spec=item.product_spec,
                                                 current_quantity=item.quantity)
                            db.session.add(inv)
                    db.session.commit()
                    flash('تم تعديل الشراء بنجاح', 'success')
            return redirect(url_for('store_transactions'))

        transaction_type = request.form.get('transaction_type')
        record_date = datetime.strptime(request.form.get('date'), '%Y-%m-%d').date()
        party_name = request.form.get('party_name')
        new_party_name = request.form.get('new_party_name', '').strip()
        if party_name == 'new' and new_party_name:
            party_name = new_party_name
        party_phone = request.form.get('party_phone')
        payment_type = request.form.get('payment_type', 'آجل')
        product_types = request.form.getlist('product_type')
        product_sizes = request.form.getlist('product_size')
        product_specs = request.form.getlist('product_spec')
        quantities = request.form.getlist('quantity')
        unit_prices = request.form.getlist('unit_price')
        new_product_types = request.form.getlist('new_product_type')
        new_product_sizes = request.form.getlist('new_product_size')
        new_product_specs = request.form.getlist('new_product_spec')
        has_items = False
        for i in range(len(product_types)):
            if product_types[i].strip() and quantities[i].strip():
                has_items = True
                break
        if not has_items:
            flash('⚠️ يجب إضافة على الأقل صنف واحد مع الكمية', 'danger')
            return redirect(url_for('store_transactions'))

        if transaction_type == 'sale':
            if party_name and not Customer.query.filter_by(name=party_name).first():
                db.session.add(Customer(name=party_name, phone=party_phone))
            new_sale = StoreSale(
                invoice_number=f"INV-{datetime.now().strftime('%Y%m%d%H%M%S')}",
                customer_name=party_name,
                customer_phone=party_phone,
                payment_type=payment_type,
                date=record_date,
                created_by=current_user.id,
                created_at=datetime.utcnow()
            )
            db.session.add(new_sale)
            for i in range(len(product_types)):
                if not product_types[i].strip() or not quantities[i].strip():
                    continue
                product_type = product_types[i]
                if product_type == 'new':
                    product_type = new_product_types[i] if i < len(new_product_types) else ''
                    if product_type and not Category.query.filter_by(name=product_type).first():
                        db.session.add(Category(name=product_type))
                product_size = product_sizes[i] if i < len(product_sizes) else ''
                if product_size == 'new':
                    product_size = new_product_sizes[i] if i < len(new_product_sizes) else ''
                    if product_size and not Size.query.filter_by(value=product_size).first():
                        db.session.add(Size(value=product_size))
                product_spec = product_specs[i] if i < len(product_specs) else ''
                if product_spec == 'new':
                    product_spec = new_product_specs[i] if i < len(new_product_specs) else ''
                    if product_spec and not Thickness.query.filter_by(value=product_spec).first():
                        db.session.add(Thickness(value=product_spec))
                item = StoreSaleItem(
                    sale_id=new_sale.id,
                    product_type=product_type,
                    product_size=product_size,
                    product_spec=product_spec,
                    quantity=float(quantities[i]) if quantities[i] else 0,
                    unit_price=float(unit_prices[i]) if unit_prices[i] else 0,
                )
                item.total = item.quantity * item.unit_price
                db.session.add(item)
                inv = StoreInventory.query.filter_by(
                    product_type=item.product_type,
                    product_size=item.product_size,
                    product_spec=item.product_spec
                ).first()
                if inv:
                    inv.current_quantity -= item.quantity
                else:
                    inv = StoreInventory(
                        product_type=item.product_type,
                        product_size=item.product_size,
                        product_spec=item.product_spec,
                        current_quantity=-item.quantity
                    )
                    db.session.add(inv)
            db.session.commit()
            total = db.session.query(db.func.sum(StoreSaleItem.total)).filter(StoreSaleItem.sale_id == new_sale.id).scalar() or 0
            db.session.add(StoreDiary(
                date=record_date,
                description=f"بيع إلى {party_name} بقيمة {total}",
                amount=total,
                created_by=current_user.id,
                created_at=datetime.utcnow()
            ))
            db.session.commit()
            flash('تم تسجيل البيع بنجاح', 'success')
        else:
            if party_name and not Supplier.query.filter_by(name=party_name).first():
                db.session.add(Supplier(name=party_name, phone=party_phone))
            new_purchase = StorePurchase(
                invoice_number=f"PUR-{datetime.now().strftime('%Y%m%d%H%M%S')}",
                supplier_name=party_name,
                supplier_phone=party_phone,
                payment_type=payment_type,
                date=record_date,
                created_by=current_user.id,
                created_at=datetime.utcnow()
            )
            db.session.add(new_purchase)
            for i in range(len(product_types)):
                if not product_types[i].strip() or not quantities[i].strip():
                    continue
                product_type = product_types[i]
                if product_type == 'new':
                    product_type = new_product_types[i] if i < len(new_product_types) else ''
                    if product_type and not Category.query.filter_by(name=product_type).first():
                        db.session.add(Category(name=product_type))
                product_size = product_sizes[i] if i < len(product_sizes) else ''
                if product_size == 'new':
                    product_size = new_product_sizes[i] if i < len(new_product_sizes) else ''
                    if product_size and not Size.query.filter_by(value=product_size).first():
                        db.session.add(Size(value=product_size))
                product_spec = product_specs[i] if i < len(product_specs) else ''
                if product_spec == 'new':
                    product_spec = new_product_specs[i] if i < len(new_product_specs) else ''
                    if product_spec and not Thickness.query.filter_by(value=product_spec).first():
                        db.session.add(Thickness(value=product_spec))
                item = StorePurchaseItem(
                    purchase_id=new_purchase.id,
                    product_type=product_type,
                    product_size=product_size,
                    product_spec=product_spec,
                    quantity=float(quantities[i]) if quantities[i] else 0,
                    unit_price=float(unit_prices[i]) if unit_prices[i] else 0,
                )
                item.total = item.quantity * item.unit_price
                db.session.add(item)
                inv = StoreInventory.query.filter_by(
                    product_type=item.product_type,
                    product_size=item.product_size,
                    product_spec=item.product_spec
                ).first()
                if inv:
                    inv.current_quantity += item.quantity
                else:
                    inv = StoreInventory(
                        product_type=item.product_type,
                        product_size=item.product_size,
                        product_spec=item.product_spec,
                        current_quantity=item.quantity
                    )
                    db.session.add(inv)
            db.session.commit()
            total = db.session.query(db.func.sum(StorePurchaseItem.total)).filter(StorePurchaseItem.purchase_id == new_purchase.id).scalar() or 0
            db.session.add(StoreDiary(
                date=record_date,
                description=f"شراء من {party_name} بقيمة {total}",
                amount=total,
                created_by=current_user.id,
                created_at=datetime.utcnow()
            ))
            db.session.commit()
            flash('تم تسجيل الشراء بنجاح', 'success')
        return redirect(url_for('store_transactions'))

    sales = StoreSale.query.order_by(StoreSale.date.asc(), StoreSale.id.asc()).all()
    purchases = StorePurchase.query.order_by(StorePurchase.date.asc(), StorePurchase.id.asc()).all()
    customers = Customer.query.order_by(Customer.name.asc()).all()
    suppliers = Supplier.query.order_by(Supplier.name.asc()).all()
    categories = Category.query.order_by(Category.name.asc()).all()
    sizes = Size.query.order_by(Size.value.asc()).all()
    thicknesses = Thickness.query.order_by(Thickness.value.asc()).all()
    return render_template('store/transactions.html',
                           sales=sales, purchases=purchases,
                           customers=customers, suppliers=suppliers,
                           categories=categories, sizes=sizes, thicknesses=thicknesses)

@app.route('/store/inventory', methods=['GET', 'POST'])
@custom_login_required
@role_required('meg', 'admin', 'mariam', 'rehab', 'ahmed', 'sayed')
def store_inventory():
    if request.method == 'POST':
        inventory_id = int(request.form.get('inventory_id'))
        item = StoreInventory.query.get_or_404(inventory_id)
        if request.form.get('delete'):
            if current_user.role in ['meg', 'admin', 'mariam', 'sayed']:
                db.session.delete(item)
                db.session.commit()
                flash('تم حذف عنصر المخزون', 'success')
        else:
            item.current_quantity = float(request.form.get('current_quantity', item.current_quantity))
            item.min_quantity = float(request.form.get('min_quantity', item.min_quantity))
            db.session.commit()
            flash('تم تحديث المخزون', 'success')
        return redirect(url_for('store_inventory'))
    inventory = StoreInventory.query.order_by(StoreInventory.product_type.asc(), StoreInventory.product_size.asc()).all()
    return render_template('store/inventory.html', inventory=inventory)

@app.route('/store/returns', methods=['GET', 'POST'])
@custom_login_required
@role_required('meg', 'admin', 'mariam', 'rehab', 'ahmed', 'sayed')
def store_returns():
    if request.method == 'POST':
        if request.form.get('delete_id'):
            record_id = int(request.form.get('delete_id'))
            record = StoreReturn.query.get_or_404(record_id)
            if can_delete_record(current_user.role, record.date):
                db.session.delete(record)
                db.session.commit()
                flash('تم حذف المرتجع بنجاح', 'success')
            return redirect(url_for('store_returns'))
        if request.form.get('edit_id'):
            record_id = int(request.form.get('edit_id'))
            record = StoreReturn.query.get_or_404(record_id)
            if can_edit(current_user.role, record.date, record.created_by, current_user.id):
                record.date = datetime.strptime(request.form.get('date'), '%Y-%m-%d').date()
                record.return_type = request.form.get('return_type')
                record.party_name = request.form.get('party_name')
                record.product_type = request.form.get('product_type')
                record.product_size = request.form.get('product_size')
                record.product_spec = request.form.get('product_spec')
                record.quantity = float(request.form.get('quantity', 0))
                record.reason = request.form.get('reason')
                db.session.commit()
                flash('تم تحديث المرتجع بنجاح', 'success')
            return redirect(url_for('store_returns'))
        record_date = datetime.strptime(request.form.get('date'), '%Y-%m-%d').date()
        return_type = request.form.get('return_type')
        party_name = request.form.get('party_name')
        product_type = request.form.get('product_type')
        product_size = request.form.get('product_size')
        product_spec = request.form.get('product_spec')
        quantity = float(request.form.get('quantity', 0))
        reason = request.form.get('reason')
        new_return = StoreReturn(date=record_date, return_type=return_type, party_name=party_name,
                                 product_type=product_type, product_size=product_size, product_spec=product_spec,
                                 quantity=quantity, reason=reason, created_by=current_user.id, created_at=datetime.utcnow())
        db.session.add(new_return)
        db.session.commit()
        flash('تم تسجيل المرتجع بنجاح', 'success')
        return redirect(url_for('store_returns'))
    returns = StoreReturn.query.order_by(StoreReturn.date.asc(), StoreReturn.id.asc()).all()
    return render_template('store/returns.html', returns=returns)

@app.route('/store/diary', methods=['GET', 'POST'])
@custom_login_required
@role_required('meg', 'admin', 'mariam', 'rehab', 'ahmed', 'sayed')
def store_diary():
    if request.method == 'POST':
        if request.form.get('delete_id'):
            record_id = int(request.form.get('delete_id'))
            record = StoreDiary.query.get_or_404(record_id)
            if can_delete_record(current_user.role, record.date):
                db.session.delete(record)
                db.session.commit()
                flash('تم حذف اليومية بنجاح', 'success')
            return redirect(url_for('store_diary'))
        if request.form.get('edit_id'):
            record_id = int(request.form.get('edit_id'))
            record = StoreDiary.query.get_or_404(record_id)
            if can_edit(current_user.role, record.date, record.created_by, current_user.id):
                record.date = datetime.strptime(request.form.get('date'), '%Y-%m-%d').date()
                record.description = request.form.get('description')
                record.amount = float(request.form.get('amount', 0))
                db.session.commit()
                flash('تم تحديث اليومية بنجاح', 'success')
            return redirect(url_for('store_diary'))
        record_date = datetime.strptime(request.form.get('date'), '%Y-%m-%d').date()
        description = request.form.get('description')
        amount = float(request.form.get('amount', 0))
        new_record = StoreDiary(date=record_date, description=description, amount=amount,
                                created_by=current_user.id, created_at=datetime.utcnow())
        db.session.add(new_record)
        db.session.commit()
        flash('تم تسجيل اليومية بنجاح', 'success')
        return redirect(url_for('store_diary'))
    diary = StoreDiary.query.order_by(StoreDiary.date.asc(), StoreDiary.id.asc()).all()
    return render_template('store/diary.html', diary=diary)
# ==================== الجرد ====================
@app.route('/inventory-audit', methods=['GET', 'POST'])
@custom_login_required
@role_required('meg', 'admin', 'mariam', 'sayed')
def inventory_audit():
    if request.method == 'POST':
        if request.form.get('delete_id'):
            record_id = int(request.form.get('delete_id'))
            record = InventoryAudit.query.get_or_404(record_id)
            db.session.delete(record)
            db.session.commit()
            flash('تم حذف سجل الجرد بنجاح', 'success')
            return redirect(url_for('inventory_audit'))
        
        audit_date = datetime.strptime(request.form.get('audit_date'), '%Y-%m-%d').date()
        audit_type = request.form.get('audit_type')
        count = 0
        
        if audit_type == 'store':
            item_names = request.form.getlist('store_item_name[]')
            item_sizes = request.form.getlist('store_item_size[]')
            item_specs = request.form.getlist('store_item_spec[]')
            system_quantities = request.form.getlist('store_system_quantity[]')
            actual_quantities = request.form.getlist('store_actual_quantity[]')
            
            for i in range(len(item_names)):
                if not item_names[i].strip() or not actual_quantities[i].strip():
                    continue
                item_name = item_names[i]
                item_size = item_sizes[i] if i < len(item_sizes) else ''
                item_spec = item_specs[i] if i < len(item_specs) else ''
                system_qty = float(system_quantities[i]) if system_quantities[i] else 0
                actual_qty = float(actual_quantities[i]) if actual_quantities[i] else 0
                difference = actual_qty - system_qty
                
                if difference < 0:
                    difference_type = 'ناقص'
                elif difference > 0:
                    difference_type = 'زيادة'
                else:
                    difference_type = 'متطابق'
                
                inv = StoreInventory.query.filter_by(
                    product_type=item_name, product_size=item_size, product_spec=item_spec
                ).first()
                if inv:
                    inv.current_quantity = actual_qty
                else:
                    inv = StoreInventory(product_type=item_name, product_size=item_size,
                                        product_spec=item_spec, current_quantity=actual_qty)
                    db.session.add(inv)
                
                if difference != 0:
                    description = f"{difference_type} جرد: {item_name} {item_size} {item_spec} - {abs(difference)}"
                    db.session.add(StoreDiary(date=audit_date, description=description, amount=0,
                                              created_by=current_user.id, created_at=datetime.utcnow()))
                
                new_record = InventoryAudit(
                    audit_date=audit_date, audit_type='store',
                    item_name=item_name, item_size=item_size, item_spec=item_spec,
                    system_quantity=system_qty, actual_quantity=actual_qty,
                    difference=difference, difference_type=difference_type,
                    created_by=current_user.id
                )
                db.session.add(new_record)
                count += 1
        
        elif audit_type == 'treasury':
            person_names = request.form.getlist('treasury_person_name[]')
            account_types = request.form.getlist('treasury_account_type[]')
            system_quantities = request.form.getlist('treasury_system_quantity[]')
            actual_quantities = request.form.getlist('treasury_actual_quantity[]')
            
            for i in range(len(person_names)):
                if not person_names[i].strip() or not actual_quantities[i].strip():
                    continue
                person_name = person_names[i]
                account_type = account_types[i] if i < len(account_types) else ''
                system_qty = float(system_quantities[i]) if system_quantities[i] else 0
                actual_qty = float(actual_quantities[i]) if actual_quantities[i] else 0
                difference = actual_qty - system_qty
                
                if difference < 0:
                    difference_type = 'ناقص'
                elif difference > 0:
                    difference_type = 'زيادة'
                else:
                    difference_type = 'متطابق'
                
                account = TreasuryAccount.query.filter_by(person_name=person_name, account_type=account_type).first()
                if account:
                    account.balance = actual_qty
                    if difference != 0:
                        txn_type = 'deposit' if difference > 0 else 'withdrawal'
                        db.session.add(TreasuryTransaction(
                            account_id=account.id, transaction_type=txn_type, amount=abs(difference),
                            source=f"{difference_type} جرد", payment_method=account_type,
                            date=audit_date, notes=f"{difference_type} جرد",
                            created_by=current_user.id, created_at=datetime.utcnow()
                        ))
                
                new_record = InventoryAudit(
                    audit_date=audit_date, audit_type='treasury',
                    item_name=person_name, account_type=account_type,
                    system_quantity=system_qty, actual_quantity=actual_qty,
                    difference=difference, difference_type=difference_type,
                    created_by=current_user.id
                )
                db.session.add(new_record)
                count += 1
        
        db.session.commit()
        flash(f'✅ تم حفظ {count} عملية جرد بنجاح', 'success')
        return redirect(url_for('inventory_audit'))
    
    records = InventoryAudit.query.order_by(InventoryAudit.audit_date.desc()).all()
    inventory = StoreInventory.query.all()
    accounts = TreasuryAccount.query.all()
    categories = Category.query.order_by(Category.name.asc()).all()
    sizes = Size.query.order_by(Size.value.asc()).all()
    thicknesses = Thickness.query.order_by(Thickness.value.asc()).all()
    return render_template('inventory_audit.html',
                           records=records, inventory=inventory, accounts=accounts,
                           categories=categories, sizes=sizes, thicknesses=thicknesses)


# ==================== Routes للإضافة السريعة ====================
@app.route('/add-quick-category', methods=['POST'])
@custom_login_required
@role_required('meg', 'admin', 'mariam', 'sayed')
def add_quick_category():
    data = request.get_json()
    name = data.get('name', '').strip()
    if not name:
        return jsonify({'success': False, 'error': 'اسم فارغ'})
    existing = Category.query.filter_by(name=name).first()
    if existing:
        return jsonify({'success': True, 'id': existing.id})
    new_cat = Category(name=name)
    db.session.add(new_cat)
    db.session.commit()
    return jsonify({'success': True, 'id': new_cat.id})


@app.route('/add-quick-size', methods=['POST'])
@custom_login_required
@role_required('meg', 'admin', 'mariam', 'sayed')
def add_quick_size():
    data = request.get_json()
    name = data.get('name', '').strip()
    if not name:
        return jsonify({'success': False, 'error': 'اسم فارغ'})
    existing = Size.query.filter_by(value=name).first()
    if existing:
        return jsonify({'success': True, 'id': existing.id})
    new_size = Size(value=name)
    db.session.add(new_size)
    db.session.commit()
    return jsonify({'success': True, 'id': new_size.id})


@app.route('/add-quick-thickness', methods=['POST'])
@custom_login_required
@role_required('meg', 'admin', 'mariam', 'sayed')
def add_quick_thickness():
    data = request.get_json()
    name = data.get('name', '').strip()
    if not name:
        return jsonify({'success': False, 'error': 'اسم فارغ'})
    existing = Thickness.query.filter_by(value=name).first()
    if existing:
        return jsonify({'success': True, 'id': existing.id})
    new_thickness = Thickness(value=name)
    db.session.add(new_thickness)
    db.session.commit()
    return jsonify({'success': True, 'id': new_thickness.id})


@app.route('/add-quick-account', methods=['POST'])
@custom_login_required
@role_required('meg', 'admin', 'mariam', 'sayed')
def add_quick_account():
    data = request.get_json()
    person_name = data.get('person_name', '').strip()
    account_type = data.get('account_type', '').strip()
    if not person_name or not account_type:
        return jsonify({'success': False, 'error': 'بيانات ناقصة'})
    existing = TreasuryAccount.query.filter_by(person_name=person_name, account_type=account_type).first()
    if existing:
        return jsonify({'success': True, 'id': existing.id})
    new_account = TreasuryAccount(person_name=person_name, account_type=account_type, balance=0)
    db.session.add(new_account)
    db.session.commit()
    return jsonify({'success': True, 'id': new_account.id})


@app.route('/add-quick-party', methods=['POST'])
@custom_login_required
@role_required('meg', 'admin', 'mariam', 'rehab', 'ahmed', 'sayed')
def add_quick_party():
    data = request.get_json()
    party_type = data.get('party_type', 'customer')
    name = data.get('name', '').strip()
    phone = data.get('phone', '').strip()
    if not name:
        return jsonify({'success': False, 'error': 'اسم فارغ'})
    if party_type == 'customer':
        existing = Customer.query.filter_by(name=name).first()
        if existing:
            return jsonify({'success': True, 'id': existing.id})
        new_party = Customer(name=name, phone=phone)
    else:
        existing = Supplier.query.filter_by(name=name).first()
        if existing:
            return jsonify({'success': True, 'id': existing.id})
        new_party = Supplier(name=name, phone=phone)
    db.session.add(new_party)
    db.session.commit()
    return jsonify({'success': True, 'id': new_party.id})


# ==================== تصدير Excel ====================
@app.route('/inventory-audit/export')
@custom_login_required
@role_required('meg', 'admin', 'mariam', 'sayed')
def inventory_audit_export():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    
    from_date_str = request.args.get('from_date')
    to_date_str = request.args.get('to_date')
    
    query = InventoryAudit.query
    if from_date_str:
        query = query.filter(InventoryAudit.audit_date >= datetime.strptime(from_date_str, '%Y-%m-%d').date())
    if to_date_str:
        query = query.filter(InventoryAudit.audit_date <= datetime.strptime(to_date_str, '%Y-%m-%d').date())
    
    records = query.order_by(InventoryAudit.audit_date.desc()).all()
    
    wb = Workbook()
    ws = wb.active
    ws.title = "الجرد"
    ws.sheet_view.rightToLeft = True
    
    headers = ['التاريخ', 'النوع', 'الصنف/الحساب', 'المقاس', 'المواصفات', 'المسجل', 'الفعلي', 'الفرق', 'الحالة']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill(start_color='7C3AED', end_color='7C3AED', fill_type='solid')
        cell.alignment = Alignment(horizontal='center')
    
    for r in records:
        ws.append([
            r.audit_date.strftime('%Y-%m-%d'),
            'مخزون' if r.audit_type == 'store' else 'خزينة',
            r.item_name or '-', r.item_size or '-', r.item_spec or '-',
            r.system_quantity, r.actual_quantity, r.difference, r.difference_type
        ])
    
    from io import BytesIO
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    return send_file(output, as_attachment=True,
                     download_name=f'audit_{datetime.now().strftime("%Y%m%d")}.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@app.route('/inventory-audit/monthly')
@custom_login_required
@role_required('meg', 'admin', 'mariam', 'sayed')
def inventory_audit_monthly():
    month = request.args.get('month', datetime.now().strftime('%Y-%m'))
    year, mon = month.split('-')
    from_date = date(int(year), int(mon), 1)
    if int(mon) == 12:
        to_date = date(int(year) + 1, 1, 1)
    else:
        to_date = date(int(year), int(mon) + 1, 1)
    
    records = InventoryAudit.query.filter(
        InventoryAudit.audit_date >= from_date,
        InventoryAudit.audit_date < to_date
    ).order_by(InventoryAudit.audit_date.desc()).all()
    
    total_shortage = sum(abs(r.difference) for r in records if r.difference < 0)
    total_surplus = sum(r.difference for r in records if r.difference > 0)
    
    return render_template('inventory_audit_monthly.html',
                           records=records, month=month,
                           total_shortage=total_shortage, total_surplus=total_surplus)


# ==================== الخزينة ====================
@app.route('/treasury')
@custom_login_required
@role_required('meg', 'admin', 'mariam', 'ahmed', 'eid', 'abdo', 'sayed')
def treasury_index():
    if current_user.role in ['meg', 'admin', 'mariam', 'sayed']:
        persons = ['الحاج أحمد', 'عيد', 'عبدالله', 'الحاج فتحي']
        for person in persons:
            for acc_type in ['كاش', 'فودافون كاش', 'انستا باي', 'شيك']:
                get_or_create_treasury_account(person, acc_type)
    else:
        person_name = current_user.full_name
        for acc_type in ['كاش', 'فودافون كاش', 'انستا باي', 'شيك']:
            get_or_create_treasury_account(person_name, acc_type)
    accounts = get_visible_accounts_for_current_user()
    if current_user.role in ['meg', 'admin', 'mariam', 'sayed']:
        transactions = TreasuryTransaction.query.order_by(TreasuryTransaction.date.asc(), TreasuryTransaction.id.asc()).limit(50).all()
    else:
        transactions = TreasuryTransaction.query.filter_by(created_by=current_user.id).order_by(TreasuryTransaction.date.asc(), TreasuryTransaction.id.asc()).limit(50).all()
    return render_template('treasury/index.html', accounts=accounts, transactions=transactions)

@app.route('/treasury/transactions', methods=['GET', 'POST'])
@custom_login_required
@role_required('meg', 'admin', 'mariam', 'ahmed', 'eid', 'abdo', 'sayed')
def treasury_transactions():
    if current_user.role in ['meg', 'admin', 'mariam', 'sayed']:
        persons = ['الحاج أحمد', 'عيد', 'عبدالله', 'الحاج فتحي']
        for person in persons:
            for acc_type in ['كاش', 'فودافون كاش', 'انستا باي', 'شيك']:
                get_or_create_treasury_account(person, acc_type)
    else:
        person_name = current_user.full_name
        for acc_type in ['كاش', 'فودافون كاش', 'انستا باي', 'شيك']:
            get_or_create_treasury_account(person_name, acc_type)
    if request.method == 'POST':
        if request.form.get('delete_id'):
            record_id = int(request.form.get('delete_id'))
            record = TreasuryTransaction.query.get_or_404(record_id)
            if current_user.role in ['meg', 'admin', 'mariam', 'sayed']:
                account = TreasuryAccount.query.get(record.account_id)
                if account:
                    if record.transaction_type == 'deposit':
                        account.balance -= record.amount
                    else:
                        account.balance += record.amount
                db.session.delete(record)
                db.session.commit()
                flash('تم حذف الحركة بنجاح', 'success')
            return redirect(url_for('treasury_transactions'))
        if request.form.get('edit_id'):
            record_id = int(request.form.get('edit_id'))
            record = TreasuryTransaction.query.get_or_404(record_id)
            if can_edit(current_user.role, record.date, record.created_by, current_user.id):
                old_account = TreasuryAccount.query.get(record.account_id)
                if old_account:
                    if record.transaction_type == 'deposit':
                        old_account.balance -= record.amount
                    else:
                        old_account.balance += record.amount
                record.date = datetime.strptime(request.form.get('date'), '%Y-%m-%d').date()
                new_account_id = int(request.form.get('account_id'))
                record.transaction_type = request.form.get('transaction_type')
                record.amount = float(request.form.get('amount', 0))
                record.source = request.form.get('source')
                record.payment_method = request.form.get('payment_method')
                record.notes = request.form.get('notes')
                record.account_id = new_account_id
                new_account = TreasuryAccount.query.get(new_account_id)
                if new_account:
                    if record.transaction_type == 'deposit':
                        new_account.balance += record.amount
                    else:
                        new_account.balance -= record.amount
                db.session.commit()
                flash('تم تحديث الحركة بنجاح', 'success')
            return redirect(url_for('treasury_transactions'))
        record_date = datetime.strptime(request.form.get('date'), '%Y-%m-%d').date()
        account_id = int(request.form.get('account_id'))
        transaction_type = request.form.get('transaction_type')
        amount = float(request.form.get('amount', 0))
        source = request.form.get('source')
        payment_method = request.form.get('payment_method')
        notes = request.form.get('notes')
        account = TreasuryAccount.query.get(account_id)
        if account:
            if transaction_type == 'deposit':
                account.balance += amount
            elif transaction_type == 'withdrawal':
                account.balance -= amount
        new_transaction = TreasuryTransaction(
            account_id=account_id, transaction_type=transaction_type, amount=amount,
            source=source, payment_method=payment_method, date=record_date, notes=notes,
            created_by=current_user.id, created_at=datetime.utcnow()
        )
        db.session.add(new_transaction)
        db.session.commit()
        flash('تم تسجيل الحركة بنجاح', 'success')
        return redirect(url_for('treasury_transactions'))
    accounts = get_visible_accounts_for_current_user()
    if current_user.role in ['meg', 'admin', 'mariam', 'sayed']:
        transactions = TreasuryTransaction.query.order_by(TreasuryTransaction.date.asc(), TreasuryTransaction.id.asc()).all()
    else:
        transactions = TreasuryTransaction.query.filter_by(created_by=current_user.id).order_by(TreasuryTransaction.date.asc(), TreasuryTransaction.id.asc()).all()
    customers = Customer.query.order_by(Customer.name.asc()).all()
    suppliers = Supplier.query.order_by(Supplier.name.asc()).all()
    return render_template('treasury/transactions.html',
                           accounts=accounts, transactions=transactions,
                           customers=customers, suppliers=suppliers)

@app.route('/treasury/transfers', methods=['GET', 'POST'])
@custom_login_required
@role_required('meg', 'admin', 'mariam', 'sayed')
def treasury_transfers():
    if request.method == 'POST':
        if request.form.get('delete_id'):
            record_id = int(request.form.get('delete_id'))
            record = TreasuryTransfer.query.get_or_404(record_id)
            if current_user.role in ['meg', 'admin', 'mariam', 'sayed']:
                db.session.delete(record)
                db.session.commit()
                flash('تم حذف التحويل بنجاح', 'success')
            return redirect(url_for('treasury_transfers'))
        if request.form.get('edit_id'):
            record_id = int(request.form.get('edit_id'))
            record = TreasuryTransfer.query.get_or_404(record_id)
            if current_user.role in ['meg', 'admin', 'mariam', 'sayed']:
                record.date = datetime.strptime(request.form.get('date'), '%Y-%m-%d').date()
                record.from_person = request.form.get('from_person')
                record.to_person = request.form.get('to_person')
                record.amount = float(request.form.get('amount', 0))
                record.payment_method = request.form.get('payment_method')
                record.notes = request.form.get('notes')
                db.session.commit()
                flash('تم تحديث التحويل بنجاح', 'success')
            return redirect(url_for('treasury_transfers'))
        date_str = request.form.get('date')
        from_person = request.form.get('from_person')
        payment_method = request.form.get('payment_method')
        notes = request.form.get('notes')
        to_persons = request.form.getlist('to_person[]')
        amounts = request.form.getlist('amount[]')
        for i in range(len(to_persons)):
            if to_persons[i].strip() and amounts[i].strip():
                transfer = TreasuryTransfer(
                    date=datetime.strptime(date_str, '%Y-%m-%d').date(),
                    from_person=from_person, to_person=to_persons[i],
                    amount=float(amounts[i]), payment_method=payment_method,
                    notes=notes, created_by=current_user.id, created_at=datetime.utcnow()
                )
                db.session.add(transfer)
        db.session.commit()
        flash('تم تسجيل التحويلات بنجاح', 'success')
        return redirect(url_for('treasury_transfers'))
    transfers = TreasuryTransfer.query.order_by(TreasuryTransfer.date.asc(), TreasuryTransfer.id.asc()).all()
    return render_template('treasury/transfers.html', transfers=transfers)

@app.route('/treasury/accounts', methods=['GET', 'POST'])
@custom_login_required
@role_required('meg', 'admin', 'mariam', 'sayed')
def treasury_accounts():
    if request.method == 'POST':
        if request.form.get('delete_id'):
            account_id = int(request.form.get('delete_id'))
            account = TreasuryAccount.query.get_or_404(account_id)
            if current_user.role in ['meg', 'admin', 'mariam', 'sayed']:
                if account.transactions:
                    flash('لا يمكن حذف حساب له معاملات', 'danger')
                else:
                    db.session.delete(account)
                    db.session.commit()
                    flash('تم حذف الحساب', 'success')
            return redirect(url_for('treasury_accounts'))
        if request.form.get('edit_id'):
            account_id = int(request.form.get('edit_id'))
            account = TreasuryAccount.query.get_or_404(account_id)
            if current_user.role in ['meg', 'admin', 'mariam', 'sayed']:
                account.person_name = request.form.get('person_name')
                account.account_type = request.form.get('account_type')
                account.balance = float(request.form.get('balance', 0))
                db.session.commit()
                flash('تم تحديث الحساب', 'success')
            return redirect(url_for('treasury_accounts'))
        person_name = request.form.get('person_name')
        account_type = request.form.get('account_type')
        balance = float(request.form.get('balance', 0))
        new_account = TreasuryAccount(person_name=person_name, account_type=account_type, balance=balance)
        db.session.add(new_account)
        db.session.commit()
        flash('تم إضافة الحساب', 'success')
        return redirect(url_for('treasury_accounts'))
    accounts = TreasuryAccount.query.order_by(TreasuryAccount.person_name.asc(), TreasuryAccount.account_type.asc()).all()
    return render_template('treasury/accounts.html', accounts=accounts)

@app.route('/financial-transactions', methods=['GET', 'POST'])
@custom_login_required
@role_required('meg', 'admin', 'mariam', 'rehab', 'ahmed', 'eid', 'abdo', 'sayed')
def financial_transactions():
    if request.method == 'POST':
        dates = request.form.getlist('date[]')
        types = request.form.getlist('transaction_type[]')
        amounts = request.form.getlist('amount[]')
        payment_methods = request.form.getlist('payment_method[]')
        notes_list = request.form.getlist('notes[]')
        from_parties = request.form.getlist('from_party[]')
        to_parties = request.form.getlist('to_party[]')
        new_from_parties = request.form.getlist('new_from_party[]')
        new_to_parties = request.form.getlist('new_to_party[]')
        for i in range(len(amounts)):
            if not amounts[i].strip():
                continue
            try:
                amount = float(amounts[i])
                record_date = datetime.strptime(dates[i], '%Y-%m-%d').date() if dates[i] else date.today()
                txn_type = types[i] if i < len(types) else 'deposit'
                payment_method = payment_methods[i] if i < len(payment_methods) else 'كاش'
                notes = notes_list[i] if i < len(notes_list) else ''
                from_party = from_parties[i] if i < len(from_parties) else ''
                to_party = to_parties[i] if i < len(to_parties) else ''
                new_from_party = new_from_parties[i] if i < len(new_from_parties) else ''
                new_to_party = new_to_parties[i] if i < len(new_to_parties) else ''
                if from_party == 'new' and new_from_party:
                    if txn_type == 'deposit':
                        if not Customer.query.filter_by(name=new_from_party).first():
                            db.session.add(Customer(name=new_from_party))
                    from_party = new_from_party
                elif from_party and from_party not in ['', 'new']:
                    if txn_type == 'deposit':
                        if not Customer.query.filter_by(name=from_party).first() and not Supplier.query.filter_by(name=from_party).first():
                            db.session.add(Customer(name=from_party))
                if to_party == 'new' and new_to_party:
                    if txn_type == 'withdrawal':
                        if not Supplier.query.filter_by(name=new_to_party).first():
                            db.session.add(Supplier(name=new_to_party))
                    to_party = new_to_party
                elif to_party and to_party not in ['', 'new']:
                    if txn_type == 'withdrawal':
                        if not Supplier.query.filter_by(name=to_party).first() and not Customer.query.filter_by(name=to_party).first():
                            db.session.add(Supplier(name=to_party))
                account = None
                if txn_type == 'deposit':
                    account_name = current_user.full_name
                    account = TreasuryAccount.query.filter_by(person_name=account_name, account_type=payment_method).first()
                    if not account:
                        account = TreasuryAccount(person_name=account_name, account_type=payment_method, balance=0)
                        db.session.add(account)
                    account.balance += amount
                    source = from_party
                    txn_type_db = 'deposit'
                elif txn_type == 'withdrawal':
                    account_name = current_user.full_name
                    account = TreasuryAccount.query.filter_by(person_name=account_name, account_type=payment_method).first()
                    if not account:
                        account = TreasuryAccount(person_name=account_name, account_type=payment_method, balance=0)
                        db.session.add(account)
                    account.balance -= amount
                    source = to_party
                    txn_type_db = 'withdrawal'
                elif txn_type == 'transfer_customer_supplier':
                    account = TreasuryAccount.query.filter_by(person_name='تحويلات العملاء', account_type='تحويل').first()
                    if not account:
                        account = TreasuryAccount(person_name='تحويلات العملاء', account_type='تحويل', balance=0)
                        db.session.add(account)
                    source = f"من {from_party} إلى {to_party}"
                    txn_type_db = 'transfer'
                new_txn = TreasuryTransaction(
                    account_id=account.id, transaction_type=txn_type_db, amount=amount,
                    source=source, payment_method=payment_method, date=record_date,
                    notes=notes, created_by=current_user.id, created_at=datetime.utcnow()
                )
                db.session.add(new_txn)
            except Exception as e:
                print(f"❌ Error: {e}")
                continue
        db.session.commit()
        flash('✅ تم تسجيل المعاملات المالية بنجاح', 'success')
        return redirect(url_for('financial_transactions'))
    transactions = TreasuryTransaction.query.order_by(TreasuryTransaction.date.asc(), TreasuryTransaction.id.asc()).all()
    accounts = TreasuryAccount.query.all()
    customers = Customer.query.order_by(Customer.name.asc()).all()
    suppliers = Supplier.query.order_by(Supplier.name.asc()).all()
    return render_template('reports/financial.html',
                           transactions=transactions, accounts=accounts,
                           customers=customers, suppliers=suppliers)

# ==================== التقارير ====================
@app.route('/reports')
@custom_login_required
@role_required('meg', 'admin', 'mariam', 'rehab', 'sayed')
def reports_index():
    return render_template('reports/index.html')

@app.route('/reports/custom')
@custom_login_required
@role_required('meg', 'admin', 'mariam', 'rehab', 'sayed')
def reports_custom():
    from_date_str = request.args.get('from_date')
    to_date_str = request.args.get('to_date')
    report_type = request.args.get('report_type', 'all')
    if not from_date_str or not to_date_str:
        return render_template('reports/custom.html',
                               from_date=None, to_date=None, report_type=report_type,
                               combined_diary=[], raw_materials=[], production=[],
                               sales=[], purchases=[], transactions=[])
    from_date = datetime.strptime(from_date_str, '%Y-%m-%d').date()
    to_date = datetime.strptime(to_date_str, '%Y-%m-%d').date()
    raw_materials = FactoryRawMaterial.query.filter(FactoryRawMaterial.date >= from_date, FactoryRawMaterial.date <= to_date).order_by(FactoryRawMaterial.date.asc()).all()
    production = FactoryProduction.query.filter(FactoryProduction.date >= from_date, FactoryProduction.date <= to_date).order_by(FactoryProduction.date.asc()).all()
    factory_diary = FactoryDiary.query.filter(FactoryDiary.date >= from_date, FactoryDiary.date <= to_date).order_by(FactoryDiary.date.asc()).all()
    sales = StoreSale.query.filter(StoreSale.date >= from_date, StoreSale.date <= to_date).order_by(StoreSale.date.asc()).all()
    purchases = StorePurchase.query.filter(StorePurchase.date >= from_date, StorePurchase.date <= to_date).order_by(StorePurchase.date.asc()).all()
    store_diary = StoreDiary.query.filter(StoreDiary.date >= from_date, StoreDiary.date <= to_date).order_by(StoreDiary.date.asc()).all()
    transactions = TreasuryTransaction.query.filter(TreasuryTransaction.date >= from_date, TreasuryTransaction.date <= to_date).order_by(TreasuryTransaction.date.asc()).all()
    combined_diary = []
    for d in factory_diary:
        combined_diary.append({'date': d.date, 'type': 'مصنع', 'description': d.description, 'amount': d.amount, 'created_by': d.created_by})
    for d in store_diary:
        combined_diary.append({'date': d.date, 'type': 'محل', 'description': d.description, 'amount': d.amount, 'created_by': d.created_by})
    combined_diary.sort(key=lambda x: x['date'])
    return render_template('reports/custom.html',
                           from_date=from_date_str, to_date=to_date_str, report_type=report_type,
                           raw_materials=raw_materials, production=production, factory_diary=factory_diary,
                           sales=sales, purchases=purchases, transactions=transactions,
                           combined_diary=combined_diary)

@app.route('/reports/customers')
@custom_login_required
@role_required('meg', 'admin', 'mariam', 'rehab', 'sayed')
def reports_customers():
    customers = Customer.query.order_by(Customer.name.asc()).all()
    customers_data = []
    for c in customers:
        sales = StoreSale.query.filter_by(customer_name=c.name).all()
        total_purchases = sum(s.total for s in sales)
        total_paid = sum(s.paid_amount for s in sales)
        customers_data.append({
            'customer': c, 'total_purchases': total_purchases,
            'total_paid': total_paid, 'remaining': total_purchases - total_paid
        })
    return render_template('reports/customers.html', customers_data=customers_data)

@app.route('/reports/customers/<int:customer_id>')
@custom_login_required
@role_required('meg', 'admin', 'mariam', 'rehab', 'sayed')
def report_single_customer(customer_id):
    customer = Customer.query.get_or_404(customer_id)
    sales = StoreSale.query.filter_by(customer_name=customer.name).order_by(StoreSale.date.asc(), StoreSale.id.asc()).all()
    total_purchases = sum(s.total for s in sales)
    total_paid = sum(s.paid_amount for s in sales)
    total_remaining = total_purchases - total_paid
    treasury_txns = TreasuryTransaction.query.filter_by(source=customer.name).order_by(TreasuryTransaction.date.asc()).all()
    return render_template('reports/customer_detail.html',
                           customer=customer, sales=sales,
                           total_purchases=total_purchases, total_paid=total_paid,
                           total_remaining=total_remaining, treasury_txns=treasury_txns)

@app.route('/reports/suppliers')
@custom_login_required
@role_required('meg', 'admin', 'mariam', 'rehab', 'sayed')
def reports_suppliers():
    suppliers = Supplier.query.order_by(Supplier.name.asc()).all()
    suppliers_data = []
    for s in suppliers:
        purchases = StorePurchase.query.filter_by(supplier_name=s.name).all()
        total_purchases = sum(p.total for p in purchases)
        total_paid = sum(p.paid_amount for p in purchases)
        suppliers_data.append({
            'supplier': s, 'total_purchases': total_purchases,
            'total_paid': total_paid, 'remaining': total_purchases - total_paid
        })
    return render_template('reports/suppliers.html', suppliers_data=suppliers_data)

@app.route('/reports/suppliers/<int:supplier_id>')
@custom_login_required
@role_required('meg', 'admin', 'mariam', 'rehab', 'sayed')
def report_single_supplier(supplier_id):
    supplier = Supplier.query.get_or_404(supplier_id)
    purchases = StorePurchase.query.filter_by(supplier_name=supplier.name).order_by(StorePurchase.date.asc(), StorePurchase.id.asc()).all()
    total_purchases = sum(p.total for p in purchases)
    total_paid = sum(p.paid_amount for p in purchases)
    total_remaining = total_purchases - total_paid
    treasury_txns = TreasuryTransaction.query.filter_by(source=supplier.name).order_by(TreasuryTransaction.date.asc()).all()
    return render_template('reports/supplier_detail.html',
                           supplier=supplier, purchases=purchases,
                           total_purchases=total_purchases, total_paid=total_paid,
                           total_remaining=total_remaining, treasury_txns=treasury_txns)

# ==================== الإدارة ====================
@app.route('/admin/users')
@custom_login_required
@role_required('meg', 'admin', 'sayed')
def admin_users():
    users = User.query.filter_by(is_hidden=False).all() if current_user.role != 'meg' else User.query.all()
    return render_template('admin/users.html', users=users)

@app.route('/admin/users/add', methods=['GET', 'POST'])
@custom_login_required
@role_required('meg', 'admin', 'sayed')
def admin_add_user():
    if request.method == 'POST':
        username = request.form.get('username')
        password = request.form.get('password')
        full_name = request.form.get('full_name')
        role = request.form.get('role')
        phone = request.form.get('phone')
        existing = User.query.filter_by(username=username).first()
        if existing:
            flash('اسم المستخدم موجود بالفعل', 'danger')
            return redirect(url_for('admin_add_user'))
        new_user = User(username=username, full_name=full_name, role=role, phone=phone,
                        is_hidden=False, created_by=current_user.id, created_at=datetime.utcnow())
        new_user.set_password(password)
        db.session.add(new_user)
        db.session.commit()
        flash('تم إضافة المستخدم بنجاح', 'success')
        return redirect(url_for('admin_users'))
    return render_template('admin/add_user.html')

@app.route('/admin/users/<int:user_id>/edit', methods=['GET', 'POST'])
@custom_login_required
@role_required('meg', 'admin', 'sayed')
def admin_edit_user(user_id):
    user = User.query.get_or_404(user_id)
    if user.role == 'meg' and current_user.role != 'meg':
        flash('غير مصرح لك بتعديل هذا المستخدم', 'danger')
        return redirect(url_for('admin_users'))
    if request.method == 'POST':
        user.full_name = request.form.get('full_name')
        user.role = request.form.get('role')
        user.phone = request.form.get('phone')
        new_password = request.form.get('password')
        if new_password:
            user.set_password(new_password)
        db.session.commit()
        flash('تم تحديث بيانات المستخدم بنجاح', 'success')
        return redirect(url_for('admin_users'))
    return render_template('admin/edit_user.html', user=user)

@app.route('/admin/users/<int:user_id>/delete', methods=['POST'])
@custom_login_required
@role_required('meg', 'admin')
def admin_delete_user(user_id):
    user = User.query.get_or_404(user_id)
    if user.role == 'meg':
        flash('لا يمكن حذف حساب MEG', 'danger')
        return redirect(url_for('admin_users'))
    db.session.delete(user)
    db.session.commit()
    flash('تم حذف المستخدم بنجاح', 'success')
    return redirect(url_for('admin_users'))

@app.route('/admin/activity')
@custom_login_required
@role_required('meg', 'admin', 'sayed')
def admin_activity():
    activities = ActivityLog.query.order_by(ActivityLog.timestamp.desc()).limit(200).all()
    inactive_users = []
    threshold = datetime.utcnow() - timedelta(days=3)
    all_users = User.query.filter_by(is_hidden=False).all()
    for user in all_users:
        if user.last_activity and user.last_activity < threshold:
            inactive_users.append(user)
    return render_template('admin/activity.html', activities=activities, inactive_users=inactive_users)

@app.route('/admin/categories')
@custom_login_required
@role_required('meg', 'admin', 'mariam', 'sayed')
def admin_categories():
    categories = Category.query.all()
    sizes = Size.query.all()
    thicknesses = Thickness.query.all()
    suppliers = Supplier.query.all()
    customers = Customer.query.all()
    return render_template('admin/categories.html',
                           categories=categories, sizes=sizes, thicknesses=thicknesses,
                           suppliers=suppliers, customers=customers)

@app.route('/admin/categories/add', methods=['POST'])
@custom_login_required
@role_required('meg', 'admin', 'mariam', 'sayed')
def admin_add_category():
    category_type = request.form.get('category_type')
    name = request.form.get('name')
    if category_type == 'category':
        if not Category.query.filter_by(name=name).first():
            db.session.add(Category(name=name))
    elif category_type == 'size':
        if not Size.query.filter_by(value=name).first():
            db.session.add(Size(value=name))
    elif category_type == 'thickness':
        if not Thickness.query.filter_by(value=name).first():
            db.session.add(Thickness(value=name))
    elif category_type == 'supplier':
        if not Supplier.query.filter_by(name=name).first():
            db.session.add(Supplier(name=name))
    elif category_type == 'customer':
        if not Customer.query.filter_by(name=name).first():
            db.session.add(Customer(name=name))
    db.session.commit()
    flash('تمت الإضافة بنجاح', 'success')
    return redirect(url_for('admin_categories'))

@app.route('/admin/categories/<string:category_type>/<int:item_id>/edit', methods=['POST'])
@custom_login_required
@role_required('meg', 'admin', 'mariam', 'sayed')
def admin_edit_category(category_type, item_id):
    new_name = request.form.get('name')
    if category_type == 'category':
        item = Category.query.get_or_404(item_id)
        item.name = new_name
    elif category_type == 'size':
        item = Size.query.get_or_404(item_id)
        item.value = new_name
    elif category_type == 'thickness':
        item = Thickness.query.get_or_404(item_id)
        item.value = new_name
    elif category_type == 'supplier':
        item = Supplier.query.get_or_404(item_id)
        item.name = new_name
        item.phone = request.form.get('phone', item.phone)
    elif category_type == 'customer':
        item = Customer.query.get_or_404(item_id)
        item.name = new_name
        item.phone = request.form.get('phone', item.phone)
    db.session.commit()
    flash('تم التعديل بنجاح', 'success')
    return redirect(url_for('admin_categories'))

@app.route('/admin/categories/<string:category_type>/<int:item_id>/delete', methods=['POST'])
@custom_login_required
@role_required('meg', 'admin', 'mariam', 'sayed')
def admin_delete_category(category_type, item_id):
    if category_type == 'category':
        item = Category.query.get_or_404(item_id)
        db.session.delete(item)
    elif category_type == 'size':
        item = Size.query.get_or_404(item_id)
        db.session.delete(item)
    elif category_type == 'thickness':
        item = Thickness.query.get_or_404(item_id)
        db.session.delete(item)
    elif category_type == 'supplier':
        item = Supplier.query.get_or_404(item_id)
        db.session.delete(item)
    elif category_type == 'customer':
        item = Customer.query.get_or_404(item_id)
        db.session.delete(item)
    db.session.commit()
    flash('تم الحذف بنجاح', 'success')
    return redirect(url_for('admin_categories'))

# ==================== الإعدادات ====================
@app.route('/settings/profile', methods=['GET', 'POST'])
@custom_login_required
def settings_profile():
    if request.method == 'POST':
        current_user.full_name = request.form.get('full_name')
        current_user.phone = request.form.get('phone')
        db.session.commit()
        flash('تم تحديث الملف الشخصي بنجاح', 'success')
        return redirect(url_for('settings_profile'))
    return render_template('settings/profile.html')

@app.route('/settings/password', methods=['GET', 'POST'])
@custom_login_required
def settings_password():
    if request.method == 'POST':
        current_password = request.form.get('current_password')
        new_password = request.form.get('new_password')
        confirm_password = request.form.get('confirm_password')
        if not current_user.check_password(current_password):
            flash('كلمة المرور الحالية غير صحيحة', 'danger')
            return redirect(url_for('settings_password'))
        if new_password != confirm_password:
            flash('كلمة المرور الجديدة غير متطابقة', 'danger')
            return redirect(url_for('settings_password'))
        if len(new_password) < 6:
            flash('كلمة المرور يجب ألا تقل عن 6 أحرف', 'danger')
            return redirect(url_for('settings_password'))
        current_user.set_password(new_password)
        db.session.commit()
        flash('تم تغيير كلمة المرور بنجاح', 'success')
        return redirect(url_for('dashboard'))
    return render_template('settings/password.html')

# ====================================================================
# ==================== الشات ====================
# ====================================================================

@app.route('/chat')
@custom_login_required
def chat_index():
    all_users = User.query.filter(
        User.is_hidden == False,
        User.id != current_user.id
    ).order_by(User.full_name.asc()).all()
    
    user_groups = ChatGroup.query.join(ChatGroupMember).filter(
        ChatGroupMember.user_id == current_user.id
    ).order_by(ChatGroup.name.asc()).all()
    
    return render_template('chat/index.html', 
                           all_users=all_users, 
                           user_groups=user_groups)


@app.route('/chat/send', methods=['POST'])
@custom_login_required
def chat_send():
    message_text = request.form.get('message', '').strip()
    chat_type = request.form.get('type')
    chat_id = request.form.get('id')
    
    if not message_text:
        return jsonify({'success': False, 'error': 'الرسالة فارغة'})
    
    try:
        if chat_type == 'user':
            recipient = User.query.get(int(chat_id))
            if not recipient:
                return jsonify({'success': False, 'error': 'المستخدم غير موجود'})
            if recipient.id == current_user.id:
                return jsonify({'success': False, 'error': 'لا يمكن إرسال رسالة لنفسك'})
            
            msg = ChatMessage(
                sender_id=current_user.id,
                recipient_id=recipient.id,
                message=message_text
            )
        elif chat_type == 'group':
            group = ChatGroup.query.get(int(chat_id))
            if not group:
                return jsonify({'success': False, 'error': 'المجموعة غير موجودة'})
            
            membership = ChatGroupMember.query.filter_by(
                group_id=group.id, user_id=current_user.id
            ).first()
            if not membership:
                return jsonify({'success': False, 'error': 'غير مصرح'})
            
            msg = ChatMessage(
                sender_id=current_user.id,
                group_id=group.id,
                message=message_text
            )
        else:
            return jsonify({'success': False, 'error': 'نوع غير معروف'})
        
        db.session.add(msg)
        db.session.commit()
        
        return jsonify({'success': True, 'id': msg.id})
    except Exception as e:
        print(f"Chat send error: {e}")
        return jsonify({'success': False, 'error': 'حدث خطأ'})


@app.route('/chat/messages/<int:user_id>')
@custom_login_required
def chat_get_messages(user_id):
    if user_id == current_user.id:
        return jsonify({'messages': []})
    
    last_id = request.args.get('last_id', 0, type=int)
    
    messages = ChatMessage.query.filter(
        ChatMessage.group_id == None,
        ChatMessage.id > last_id,
        or_(
            and_(ChatMessage.sender_id == current_user.id, ChatMessage.recipient_id == user_id),
            and_(ChatMessage.sender_id == user_id, ChatMessage.recipient_id == current_user.id)
        )
    ).order_by(ChatMessage.created_at.asc()).all()
    
    ChatMessage.query.filter(
        ChatMessage.sender_id == user_id,
        ChatMessage.recipient_id == current_user.id,
        ChatMessage.is_read == False
    ).update({'is_read': True})
    db.session.commit()
    
    result = []
    for msg in messages:
        result.append({
            'id': msg.id,
            'sender_id': msg.sender_id,
            'sender_name': msg.sender.full_name if msg.sender else '',
            'message': msg.display_message,
            'is_deleted': msg.is_deleted,
            'time': msg.created_at.strftime('%H:%M'),
            'is_read': msg.is_read
        })
    
    return jsonify({'messages': result})


@app.route('/chat/group-messages/<int:group_id>')
@custom_login_required
def chat_get_group_messages(group_id):
    membership = ChatGroupMember.query.filter_by(
        group_id=group_id, user_id=current_user.id
    ).first()
    if not membership:
        return jsonify({'messages': []})
    
    last_id = request.args.get('last_id', 0, type=int)
    
    messages = ChatMessage.query.filter(
        ChatMessage.group_id == group_id,
        ChatMessage.id > last_id
    ).order_by(ChatMessage.created_at.asc()).all()
    
    ChatMessage.query.filter(
        ChatMessage.group_id == group_id,
        ChatMessage.sender_id != current_user.id,
        ChatMessage.is_read == False
    ).update({'is_read': True})
    db.session.commit()
    
    result = []
    for msg in messages:
        result.append({
            'id': msg.id,
            'sender_id': msg.sender_id,
            'sender_name': msg.sender.full_name if msg.sender else '',
            'message': msg.display_message,
            'is_deleted': msg.is_deleted,
            'time': msg.created_at.strftime('%H:%M'),
            'is_read': msg.is_read
        })
    
    return jsonify({'messages': result})


@app.route('/chat/delete-message/<int:msg_id>', methods=['POST'])
@custom_login_required
def chat_delete_message(msg_id):
    msg = ChatMessage.query.get_or_404(msg_id)
    
    if msg.sender_id != current_user.id:
        return jsonify({'success': False, 'error': 'غير مصرح'})
    
    if msg.is_deleted:
        return jsonify({'success': True})
    
    msg.is_deleted = True
    db.session.commit()
    
    return jsonify({'success': True})


@app.route('/chat/create-group', methods=['POST'])
@custom_login_required
@role_required('meg', 'admin', 'sayed')
def chat_create_group():
    name = request.form.get('name', '').strip()
    member_ids = request.form.getlist('members[]')
    
    if not name:
        flash('اسم المجموعة مطلوب', 'danger')
        return redirect(url_for('chat_index'))
    
    group = ChatGroup(name=name, created_by=current_user.id)
    db.session.add(group)
    db.session.flush()
    
    db.session.add(ChatGroupMember(group_id=group.id, user_id=current_user.id))
    
    for uid in member_ids:
        try:
            uid = int(uid)
            if uid != current_user.id:
                if User.query.get(uid):
                    db.session.add(ChatGroupMember(group_id=group.id, user_id=uid))
        except:
            continue
    
    db.session.commit()
    flash('تم إنشاء المجموعة بنجاح', 'success')
    return redirect(url_for('chat_index'))


@app.route('/chat/unread-count')
@custom_login_required
def chat_unread_count():
    by_user = {}
    unread_users = db.session.query(
        ChatMessage.sender_id,
        db.func.count(ChatMessage.id)
    ).filter(
        ChatMessage.recipient_id == current_user.id,
        ChatMessage.is_read == False,
        ChatMessage.is_deleted == False,
        ChatMessage.group_id == None
    ).group_by(ChatMessage.sender_id).all()
    
    for sender_id, count in unread_users:
        by_user[str(sender_id)] = count
    
    total = sum(by_user.values())
    
    return jsonify({'total': total, 'by_user': by_user, 'by_group': {}})    
# ====================================================================
# ==================== شركة الماسة ====================
# ====================================================================

# --------------------------------------------------------------------
# الصفحة الرئيسية
# --------------------------------------------------------------------
@app.route('/almasa')
@custom_login_required
@role_required('sayed', 'dina', 'admin', 'meg')
def almasa_index():
    partnership_cranes = AlMasaCrane.query.filter_by(crane_type='partnership').all()
    private_cranes = AlMasaPrivateCrane.query.all()
    supplies = AlMasaSupply.query.all()
    
    # الإجماليات
    total_partnership = sum(
        sum(o.rental_total or 0 for o in c.operations) 
        for c in partnership_cranes
    )
    total_private = sum(
        sum(o.rental_total or 0 for o in c.operations) 
        for c in private_cranes
    )
    total_supply = sum(
        sum(o.rental_total or 0 for o in s.operations) 
        for s in supplies
    )
    grand_total = total_partnership + total_private + total_supply
    
    # غير مدفوع
    total_unpaid = 0
    total_unpaid += sum(
        o.rental_total or 0 
        for o in AlMasaOperation.query.filter(
            (AlMasaOperation.check_received == False) | (AlMasaOperation.check_received == None)
        ).all()
    )
    total_unpaid += sum(
        o.rental_total or 0 
        for o in AlMasaPrivateOperation.query.filter(
            (AlMasaPrivateOperation.check_received == False) | (AlMasaPrivateOperation.check_received == None)
        ).all()
    )
    total_unpaid += sum(
        o.rental_total or 0 
        for o in AlMasaSupplyOperation.query.filter(
            (AlMasaSupplyOperation.check_received == False) | (AlMasaSupplyOperation.check_received == None)
        ).all()
    )
    
    # مدفوع
    total_paid = grand_total - total_unpaid
    
    # إجمالي الأرباح
    total_profit = 0
    # أوناش مشاركة
    for crane in partnership_cranes:
        total_rental_c = sum(o.rental_total or 0 for o in crane.operations)
        total_supply_c = sum(o.supply_total or 0 for o in crane.operations)
        total_expenses_c = sum(e.amount or 0 for e in crane.expenses)
        tax_85_c = sum(
            (o.supply_total * o.tax_85_value / 100) 
            for o in crane.operations if o.tax_85_enabled
        )
        admin = total_rental_c - total_supply_c
        supply = total_supply_c - total_expenses_c - tax_85_c
        total_profit += admin + supply
    
    # أوناش خاصة
    for crane in private_cranes:
        total_rental_c = sum(o.rental_total or 0 for o in crane.operations)
        total_expenses_c = sum(e.amount or 0 for e in crane.expenses)
        tax_85_c = sum(
            (o.rental_total * o.tax_85_value / 100) 
            for o in crane.operations if o.tax_85_enabled
        )
        total_profit += total_rental_c - total_expenses_c - tax_85_c
    
    # توريدات
    for supply in supplies:
        total_rental_s = sum(o.rental_total or 0 for o in supply.operations)
        total_supply_s = sum(o.supply_total or 0 for o in supply.operations)
        total_expenses_s = sum(e.amount or 0 for e in supply.expenses)
        tax_85_s = sum(
            (o.rental_total * o.tax_85_value / 100) 
            for o in supply.operations if o.tax_85_enabled
        )
        total_profit += total_rental_s - total_supply_s - total_expenses_s - tax_85_s
    
    return render_template('almasa/index.html',
                           partnership_cranes=partnership_cranes,
                           private_cranes=private_cranes,
                           supplies=supplies,
                           total_partnership=total_partnership,
                           total_private=total_private,
                           total_supply=total_supply,
                           grand_total=grand_total,
                           total_unpaid=total_unpaid,
                           total_paid=total_paid,
                           total_profit=total_profit)


# ====================================================================
# ==================== النوع 1: ونش مشاركة ====================
# ====================================================================
@app.route('/almasa/cranes', methods=['GET', 'POST'])
@custom_login_required
@role_required('sayed', 'dina', 'admin', 'meg')
def almasa_cranes():
    if request.method == 'POST':
        name = request.form.get('name')
        notes = request.form.get('notes')
        partner_names = request.form.getlist('partner_name[]')
        partner_percentages = request.form.getlist('partner_percentage[]')
        is_basic = request.form.getlist('is_basic[]')
        new_crane = AlMasaCrane(name=name, notes=notes, crane_type='partnership')
        db.session.add(new_crane)
        db.session.flush()
        for i in range(len(partner_names)):
            if partner_names[i].strip():
                partner = AlMasaPartner(
                    crane_id=new_crane.id, name=partner_names[i],
                    percentage=float(partner_percentages[i]) if partner_percentages[i] else 0,
                    is_basic=True if i < len(is_basic) and is_basic[i] == '1' else False
                )
                db.session.add(partner)
        db.session.commit()
        flash('تم إضافة الونش بنجاح', 'success')
        return redirect(url_for('almasa_cranes'))
    cranes = AlMasaCrane.query.filter_by(crane_type='partnership').order_by(AlMasaCrane.id.asc()).all()
    return render_template('almasa/cranes.html', cranes=cranes)


@app.route('/almasa/cranes/<int:crane_id>')
@custom_login_required
@role_required('sayed', 'dina', 'admin', 'meg')
def almasa_crane_detail(crane_id):
    crane = AlMasaCrane.query.get_or_404(crane_id)
    operations = AlMasaOperation.query.filter_by(crane_id=crane_id).order_by(AlMasaOperation.start_date.asc()).all()
    checks = AlMasaCheck.query.filter_by(crane_id=crane_id).order_by(AlMasaCheck.due_date.asc()).all()
    expenses = AlMasaExpense.query.filter_by(crane_id=crane_id).order_by(AlMasaExpense.date.asc()).all()
    partners = AlMasaPartner.query.filter_by(crane_id=crane_id).all()
    
    total_rental = sum(o.rental_total or 0 for o in operations)
    total_supply = sum(o.supply_total or 0 for o in operations)
    total_expenses = sum(e.amount or 0 for e in expenses)
    total_paid = sum(o.rental_total or 0 for o in operations if o.check_received)
    total_unpaid = sum(o.rental_total or 0 for o in operations if not o.check_received)
    total_operations = total_rental
    
    return render_template('almasa/crane_detail.html',
                           crane=crane, operations=operations, checks=checks,
                           expenses=expenses, partners=partners,
                           total_rental=total_rental, total_supply=total_supply,
                           total_expenses=total_expenses,
                           total_paid=total_paid, total_unpaid=total_unpaid,
                           total_operations=total_operations)


@app.route('/almasa/cranes/<int:crane_id>/unpaid')
@custom_login_required
@role_required('sayed', 'dina', 'admin', 'meg')
def almasa_crane_unpaid(crane_id):
    crane = AlMasaCrane.query.get_or_404(crane_id)
    operations = AlMasaOperation.query.filter_by(crane_id=crane_id).filter(
        (AlMasaOperation.check_received == False) | (AlMasaOperation.check_received == None)
    ).order_by(AlMasaOperation.start_date.asc()).all()
    total_unpaid = sum(o.rental_total or 0 for o in operations)
    return render_template('almasa/crane_unpaid.html',
                           crane=crane, operations=operations, total_unpaid=total_unpaid)


@app.route('/almasa/cranes/<int:crane_id>/delete', methods=['POST'])
@custom_login_required
@role_required('sayed', 'dina', 'admin', 'meg')
def almasa_crane_delete(crane_id):
    crane = AlMasaCrane.query.get_or_404(crane_id)
    db.session.delete(crane)
    db.session.commit()
    flash('تم حذف الونش بنجاح', 'success')
    return redirect(url_for('almasa_cranes'))


# ====================================================================
# ✅✅✅ almasa_add_operation — الحسابات الجديدة
# ====================================================================
@app.route('/almasa/operations/add', methods=['POST'])
@custom_login_required
@role_required('sayed', 'dina', 'admin', 'meg')
def almasa_add_operation():
    crane_id = int(request.form.get('crane_id'))
    start_date = datetime.strptime(request.form.get('start_date'), '%Y-%m-%d').date()
    end_date = datetime.strptime(request.form.get('end_date'), '%Y-%m-%d').date() if request.form.get('end_date') else None
    
    # القيم اليومية
    supply_value = float(request.form.get('supply_value', 0))
    rental_value = float(request.form.get('rental_value', 0))
    
    # ✅ الحقول الجديدة — عدد الأيام (منفصلة)
    rental_days = float(request.form.get('rental_days', 0))
    supply_days = float(request.form.get('supply_days', 0))
    
    # ✅ الحقول الجديدة — ساعات إضافية (منفصلة)
    rental_extra_hours = float(request.form.get('rental_extra_hours', 0))
    rental_hour_rate = float(request.form.get('rental_hour_rate', 0))
    supply_extra_hours = float(request.form.get('supply_extra_hours', 0))
    supply_hour_rate = float(request.form.get('supply_hour_rate', 0))
    
    # ✅ الحقول الجديدة — أيام الطريق (منفصلة)
    rental_travel_days = float(request.form.get('rental_travel_days', 0))
    supply_travel_days = float(request.form.get('supply_travel_days', 0))
    
    # ✅ الحقول الجديدة — سعر يوم الطريق (منفصلة)
    travel_rental_rate = float(request.form.get('travel_rental_rate', 0))
    travel_supply_rate = float(request.form.get('travel_supply_rate', 0))
    
    # ✅ الحقول الجديدة — الدفع
    payment_check_number = request.form.get('payment_check_number', '').strip()
    payment_check_due_date_str = request.form.get('payment_check_due_date')
    payment_check_due_date = datetime.strptime(payment_check_due_date_str, '%Y-%m-%d').date() if payment_check_due_date_str else None
    payment_account_name = request.form.get('payment_account_name', '').strip()
    payment_account_number = request.form.get('payment_account_number', '').strip()
    payment_date_str = request.form.get('payment_date')
    payment_date = datetime.strptime(payment_date_str, '%Y-%m-%d').date() if payment_date_str else None
    
    # الضرايب
    tax_14_enabled = request.form.get('tax_14_enabled') == 'on'
    tax_14_value = float(request.form.get('tax_14_value', 14))
    tax_85_enabled = request.form.get('tax_85_enabled') == 'on'
    tax_85_value = float(request.form.get('tax_85_value', 8.5))
    
    # الفاتورة
    invoice_number = request.form.get('invoice_number', '').strip()
    invoice_date_str = request.form.get('invoice_date')
    invoice_date = datetime.strptime(invoice_date_str, '%Y-%m-%d').date() if invoice_date_str else None
    project_name = request.form.get('project_name', '').strip()
    project_location = request.form.get('project_location', '').strip()
    payment_method = request.form.get('payment_method', '').strip()
    
    # ✅ الحسابات الجديدة
    rental_extra_value = (rental_extra_hours * rental_hour_rate) + (rental_travel_days * travel_rental_rate)
    supply_extra_value = (supply_extra_hours * supply_hour_rate) + (supply_travel_days * travel_supply_rate)
    rental_total = (rental_value * rental_days) + rental_extra_value
    supply_total = (supply_value * supply_days) + supply_extra_value
    
    operation = AlMasaOperation(
        crane_id=crane_id, start_date=start_date, end_date=end_date,
        supply_value=supply_value, rental_value=rental_value,
        supply_total=supply_total, rental_total=rental_total,
        rental_days=rental_days, supply_days=supply_days,
        rental_extra_hours=rental_extra_hours, rental_hour_rate=rental_hour_rate,
        supply_extra_hours=supply_extra_hours, supply_hour_rate=supply_hour_rate,
        rental_travel_days=rental_travel_days, supply_travel_days=supply_travel_days,
        travel_rental_rate=travel_rental_rate, travel_supply_rate=travel_supply_rate,
        payment_check_number=payment_check_number,
        payment_check_due_date=payment_check_due_date,
        payment_account_name=payment_account_name,
        payment_account_number=payment_account_number,
        payment_date=payment_date,
        tax_14_enabled=tax_14_enabled, tax_14_value=tax_14_value,
        tax_85_enabled=tax_85_enabled, tax_85_value=tax_85_value,
        invoice_number=invoice_number, invoice_date=invoice_date,
        project_name=project_name, project_location=project_location,
        payment_method=payment_method,
        created_by=current_user.id
    )
    db.session.add(operation)
    db.session.commit()
    flash('تم إضافة العملية بنجاح', 'success')
    return redirect(url_for('almasa_crane_detail', crane_id=crane_id))


@app.route('/almasa/operations/<int:op_id>/delete', methods=['POST'])
@custom_login_required
@role_required('sayed', 'dina', 'admin', 'meg')
def almasa_operation_delete(op_id):
    operation = AlMasaOperation.query.get_or_404(op_id)
    crane_id = operation.crane_id
    AlMasaCheckOperation.query.filter_by(operation_id=op_id).delete()
    db.session.delete(operation)
    db.session.commit()
    flash('تم حذف العملية بنجاح', 'success')
    return redirect(url_for('almasa_crane_detail', crane_id=crane_id))


# ====================================================================
# ✅✅✅ almasa_operation_edit — الحسابات الجديدة
# ====================================================================
@app.route('/almasa/operations/<int:op_id>/edit', methods=['POST'])
@custom_login_required
@role_required('sayed', 'dina', 'admin', 'meg')
def almasa_operation_edit(op_id):
    operation = AlMasaOperation.query.get_or_404(op_id)
    operation.start_date = datetime.strptime(request.form.get('start_date'), '%Y-%m-%d').date()
    operation.end_date = datetime.strptime(request.form.get('end_date'), '%Y-%m-%d').date() if request.form.get('end_date') else None
    
    # القيم اليومية
    operation.supply_value = float(request.form.get('supply_value', 0))
    operation.rental_value = float(request.form.get('rental_value', 0))
    
    # ✅ الحقول الجديدة
    operation.rental_days = float(request.form.get('rental_days', 0))
    operation.supply_days = float(request.form.get('supply_days', 0))
    operation.rental_extra_hours = float(request.form.get('rental_extra_hours', 0))
    operation.rental_hour_rate = float(request.form.get('rental_hour_rate', 0))
    operation.supply_extra_hours = float(request.form.get('supply_extra_hours', 0))
    operation.supply_hour_rate = float(request.form.get('supply_hour_rate', 0))
    operation.rental_travel_days = float(request.form.get('rental_travel_days', 0))
    operation.supply_travel_days = float(request.form.get('supply_travel_days', 0))
    operation.travel_rental_rate = float(request.form.get('travel_rental_rate', 0))
    operation.travel_supply_rate = float(request.form.get('travel_supply_rate', 0))
    
    # الدفع
    operation.payment_check_number = request.form.get('payment_check_number', '').strip()
    payment_check_due_date_str = request.form.get('payment_check_due_date')
    operation.payment_check_due_date = datetime.strptime(payment_check_due_date_str, '%Y-%m-%d').date() if payment_check_due_date_str else None
    operation.payment_account_name = request.form.get('payment_account_name', '').strip()
    operation.payment_account_number = request.form.get('payment_account_number', '').strip()
    payment_date_str = request.form.get('payment_date')
    operation.payment_date = datetime.strptime(payment_date_str, '%Y-%m-%d').date() if payment_date_str else None
    
    # الضرايب
    operation.tax_14_enabled = request.form.get('tax_14_enabled') == 'on'
    operation.tax_14_value = float(request.form.get('tax_14_value', 14))
    operation.tax_85_enabled = request.form.get('tax_85_enabled') == 'on'
    operation.tax_85_value = float(request.form.get('tax_85_value', 8.5))
    
    # الفاتورة
    operation.invoice_number = request.form.get('invoice_number', '').strip()
    invoice_date_str = request.form.get('invoice_date')
    operation.invoice_date = datetime.strptime(invoice_date_str, '%Y-%m-%d').date() if invoice_date_str else None
    operation.project_name = request.form.get('project_name', '').strip()
    operation.project_location = request.form.get('project_location', '').strip()
    operation.payment_method = request.form.get('payment_method', '').strip()
    
    # ✅ الحسابات الجديدة
    rental_extra_value = (operation.rental_extra_hours * operation.rental_hour_rate) + (operation.rental_travel_days * operation.travel_rental_rate)
    supply_extra_value = (operation.supply_extra_hours * operation.supply_hour_rate) + (operation.supply_travel_days * operation.travel_supply_rate)
    operation.rental_total = (operation.rental_value * operation.rental_days) + rental_extra_value
    operation.supply_total = (operation.supply_value * operation.supply_days) + supply_extra_value
    
    db.session.commit()
    flash('تم تعديل العملية بنجاح', 'success')
    return redirect(url_for('almasa_crane_detail', crane_id=operation.crane_id))


@app.route('/almasa/operations/<int:op_id>/update-notes', methods=['POST'])
@custom_login_required
@role_required('sayed', 'dina', 'admin', 'meg')
def almasa_operation_update_notes(op_id):
    operation = AlMasaOperation.query.get_or_404(op_id)
    operation.notes = request.form.get('notes', '')
    db.session.commit()
    flash('تم تحديث الملاحظات بنجاح', 'success')
    return redirect(url_for('almasa_crane_detail', crane_id=operation.crane_id))


@app.route('/almasa/operations/<int:op_id>/receive-check', methods=['POST'])
@custom_login_required
@role_required('sayed', 'dina', 'admin', 'meg')
def almasa_operation_receive_check(op_id):
    operation = AlMasaOperation.query.get_or_404(op_id)
    check_number = request.form.get('check_number', '').strip()
    issue_date_str = request.form.get('issue_date')
    issue_date = datetime.strptime(issue_date_str, '%Y-%m-%d').date() if issue_date_str else date.today()
    check = AlMasaCheck(
        crane_id=operation.crane_id, operation_id=op_id,
        check_number=check_number, company_name=operation.project_name,
        amount = operation.supply_total + (operation.supply_total * (operation.tax_14_value / 100) if operation.tax_14_enabled else 0)
    )
    db.session.add(check)
    operation.check_received = True
    operation.check_received_date = issue_date
    db.session.commit()
    flash('تم تسجيل استلام الشيك بنجاح', 'success')
    return redirect(url_for('almasa_crane_detail', crane_id=operation.crane_id))


@app.route('/almasa/expenses/add', methods=['POST'])
@custom_login_required
@role_required('sayed', 'dina', 'admin', 'meg')
def almasa_add_expense():
    crane_id = int(request.form.get('crane_id'))
    operation_id = request.form.get('operation_id')
    operation_id = int(operation_id) if operation_id else None
    date_str = datetime.strptime(request.form.get('date'), '%Y-%m-%d').date()
    expense_type = request.form.get('expense_type')
    amount = float(request.form.get('amount', 0))
    paid_by = request.form.get('paid_by', '').strip()
    notes = request.form.get('notes')
    expense = AlMasaExpense(crane_id=crane_id, operation_id=operation_id, date=date_str,
                            expense_type=expense_type, amount=amount, paid_by=paid_by,
                            notes=notes, created_by=current_user.id)
    db.session.add(expense)
    db.session.commit()
    flash('تم إضافة المصروف بنجاح', 'success')
    return redirect(url_for('almasa_crane_detail', crane_id=crane_id))


@app.route('/almasa/expenses/<int:exp_id>/delete', methods=['POST'])
@custom_login_required
@role_required('sayed', 'dina', 'admin', 'meg')
def almasa_expense_delete(exp_id):
    expense = AlMasaExpense.query.get_or_404(exp_id)
    crane_id = expense.crane_id
    db.session.delete(expense)
    db.session.commit()
    flash('تم حذف المصروف بنجاح', 'success')
    return redirect(url_for('almasa_crane_detail', crane_id=crane_id))


@app.route('/almasa/expenses/<int:exp_id>/edit', methods=['POST'])
@custom_login_required
@role_required('sayed', 'dina', 'admin', 'meg')
def almasa_expense_edit(exp_id):
    expense = AlMasaExpense.query.get_or_404(exp_id)
    expense.date = datetime.strptime(request.form.get('date'), '%Y-%m-%d').date()
    expense.expense_type = request.form.get('expense_type')
    expense.amount = float(request.form.get('amount', 0))
    expense.paid_by = request.form.get('paid_by', '').strip()
    expense.notes = request.form.get('notes')
    db.session.commit()
    flash('تم تعديل المصروف بنجاح', 'success')
    return redirect(url_for('almasa_crane_detail', crane_id=expense.crane_id))


@app.route('/almasa/checks/<int:check_id>/delete', methods=['POST'])
@custom_login_required
@role_required('sayed', 'dina', 'admin', 'meg')
def almasa_check_delete(check_id):
    check = AlMasaCheck.query.get_or_404(check_id)
    crane_id = check.crane_id
    AlMasaCheckOperation.query.filter_by(check_id=check_id).delete()
    db.session.delete(check)
    db.session.commit()
    flash('تم حذف الشيك بنجاح', 'success')
    return redirect(url_for('almasa_crane_detail', crane_id=crane_id))


@app.route('/almasa/checks/<int:check_id>/report')
@custom_login_required
@role_required('sayed', 'dina', 'admin', 'meg')
def almasa_check_report(check_id):
    check = AlMasaCheck.query.get_or_404(check_id)
    crane = AlMasaCrane.query.get(check.crane_id)
    check_operations = AlMasaCheckOperation.query.filter_by(check_id=check_id).all()
    operations_list = []
    total_rental = 0
    total_supply = 0
    total_days = 0
    for co in check_operations:
        operation = AlMasaOperation.query.get(co.operation_id)
        if operation:
            operations_list.append(operation)
            total_rental += operation.rental_total
            total_supply += operation.supply_total
            total_days += operation.days_count
    if operations_list:
        min_date = min(o.start_date for o in operations_list)
        dates_with_end = [o.end_date for o in operations_list if o.end_date]
        max_date = max(dates_with_end) if dates_with_end else max(o.start_date for o in operations_list)
        expenses = AlMasaExpense.query.filter(
            AlMasaExpense.crane_id == check.crane_id,
            AlMasaExpense.date >= min_date,
            AlMasaExpense.date <= max_date
        ).all()
    else:
        expenses = []
    total_expenses = sum(e.amount for e in expenses)
    partners = AlMasaPartner.query.filter_by(crane_id=check.crane_id).all()
    non_basic = [p for p in partners if not p.is_basic]
    if len(non_basic) == 0:
        total_supply = total_rental
    first_op = operations_list[0] if operations_list else None
    tax_14_value = first_op.tax_14_value if first_op and first_op.tax_14_enabled else 0
    tax_85_value = first_op.tax_85_value if first_op and first_op.tax_85_enabled else 0
    check_amount = total_supply + (total_supply * (tax_14_value / 100))
    tax_14 = total_supply * (tax_14_value / 100)
    admin_profit = total_rental - total_supply
    tax_85_amount = total_supply * (tax_85_value / 100)
    supply_profit = total_supply - total_expenses - tax_85_amount
    basic_partners = [p for p in partners if p.is_basic]
    basic_count = len(basic_partners) if basic_partners else 1
    basic_diff = admin_profit / basic_count
    partners_profit = []
    for p in partners:
        if p.is_basic:
            actual_share = basic_diff
            default_share = supply_profit * (p.percentage / 100)
        else:
            actual_share = 0
            default_share = supply_profit * (p.percentage / 100)
        partners_profit.append({
            'partner': p, 'actual_share': actual_share,
            'default_share': default_share, 'total_share': actual_share + default_share
        })
    return render_template('almasa/check_report.html',
                           check=check, crane=crane,
                           operations=operations_list, expenses=expenses,
                           total_rental=total_rental, total_supply=total_supply,
                           total_days=total_days, total_expenses=total_expenses,
                           check_amount=check_amount, tax_14=tax_14,
                           tax_85_amount=tax_85_amount,
                           admin_profit=admin_profit, supply_profit=supply_profit,
                           basic_diff=basic_diff,
                           partners=partners, partners_profit=partners_profit,
                           report_type='partnership')    
# ====================================================================
# ==================== النوع 2: ونش خاص ====================
# ====================================================================
@app.route('/almasa/private-cranes', methods=['GET', 'POST'])
@custom_login_required
@role_required('sayed', 'dina', 'admin', 'meg')
def almasa_private_cranes():
    if request.method == 'POST':
        name = request.form.get('name')
        owner_name = request.form.get('owner_name')
        notes = request.form.get('notes')
        new_crane = AlMasaPrivateCrane(name=name, owner_name=owner_name, notes=notes)
        db.session.add(new_crane)
        db.session.commit()
        flash('تم إضافة الونش الخاص بنجاح', 'success')
        return redirect(url_for('almasa_private_cranes'))
    cranes = AlMasaPrivateCrane.query.order_by(AlMasaPrivateCrane.id.asc()).all()
    return render_template('almasa/private_cranes.html', cranes=cranes)


@app.route('/almasa/private-cranes/<int:crane_id>')
@custom_login_required
@role_required('sayed', 'dina', 'admin', 'meg')
def almasa_private_crane_detail(crane_id):
    crane = AlMasaPrivateCrane.query.get_or_404(crane_id)
    operations = AlMasaPrivateOperation.query.filter_by(crane_id=crane_id).order_by(AlMasaPrivateOperation.start_date.asc()).all()
    checks = AlMasaPrivateCheck.query.filter_by(crane_id=crane_id).order_by(AlMasaPrivateCheck.due_date.asc()).all()
    expenses = AlMasaPrivateExpense.query.filter_by(crane_id=crane_id).order_by(AlMasaPrivateExpense.date.asc()).all()
    total_rental = sum(o.rental_total or 0 for o in operations)
    total_supply = sum(o.supply_total or 0 for o in operations)
    total_expenses = sum(e.amount or 0 for e in expenses)
    total_paid = sum(o.rental_total or 0 for o in operations if o.check_received)
    total_unpaid = sum(o.rental_total or 0 for o in operations if not o.check_received)
    total_operations = total_rental
    return render_template('almasa/private_crane_detail.html',
                           crane=crane, operations=operations,
                           checks=checks, expenses=expenses,
                           total_rental=total_rental, total_supply=total_supply,
                           total_expenses=total_expenses,
                           total_paid=total_paid, total_unpaid=total_unpaid,
                           total_operations=total_operations)


@app.route('/almasa/private-cranes/<int:crane_id>/unpaid')
@custom_login_required
@role_required('sayed', 'dina', 'admin', 'meg')
def almasa_private_crane_unpaid(crane_id):
    crane = AlMasaPrivateCrane.query.get_or_404(crane_id)
    operations = AlMasaPrivateOperation.query.filter_by(crane_id=crane_id).filter(
        (AlMasaPrivateOperation.check_received == False) | (AlMasaPrivateOperation.check_received == None)
    ).order_by(AlMasaPrivateOperation.start_date.asc()).all()
    total_unpaid = sum(o.rental_total or 0 for o in operations)
    return render_template('almasa/private_crane_unpaid.html',
                           crane=crane, operations=operations, total_unpaid=total_unpaid)


@app.route('/almasa/private-cranes/<int:crane_id>/delete', methods=['POST'])
@custom_login_required
@role_required('sayed', 'dina', 'admin', 'meg')
def almasa_private_crane_delete(crane_id):
    crane = AlMasaPrivateCrane.query.get_or_404(crane_id)
    db.session.delete(crane)
    db.session.commit()
    flash('تم حذف الونش الخاص بنجاح', 'success')
    return redirect(url_for('almasa_private_cranes'))


# ====================================================================
# ✅✅✅ almasa_add_private_operation — الحسابات الجديدة
# ====================================================================
@app.route('/almasa/private-operations/add', methods=['POST'])
@custom_login_required
@role_required('sayed', 'dina', 'admin', 'meg')
def almasa_add_private_operation():
    crane_id = int(request.form.get('crane_id'))
    start_date = datetime.strptime(request.form.get('start_date'), '%Y-%m-%d').date()
    end_date = datetime.strptime(request.form.get('end_date'), '%Y-%m-%d').date() if request.form.get('end_date') else None
    
    supply_value = float(request.form.get('supply_value', 0))
    rental_value = float(request.form.get('rental_value', 0))
    
    rental_days = float(request.form.get('rental_days', 0))
    supply_days = float(request.form.get('supply_days', 0))
    rental_extra_hours = float(request.form.get('rental_extra_hours', 0))
    rental_hour_rate = float(request.form.get('rental_hour_rate', 0))
    supply_extra_hours = float(request.form.get('supply_extra_hours', 0))
    supply_hour_rate = float(request.form.get('supply_hour_rate', 0))
    rental_travel_days = float(request.form.get('rental_travel_days', 0))
    supply_travel_days = float(request.form.get('supply_travel_days', 0))
    travel_rental_rate = float(request.form.get('travel_rental_rate', 0))
    travel_supply_rate = float(request.form.get('travel_supply_rate', 0))
    
    payment_check_number = request.form.get('payment_check_number', '').strip()
    payment_check_due_date_str = request.form.get('payment_check_due_date')
    payment_check_due_date = datetime.strptime(payment_check_due_date_str, '%Y-%m-%d').date() if payment_check_due_date_str else None
    payment_account_name = request.form.get('payment_account_name', '').strip()
    payment_account_number = request.form.get('payment_account_number', '').strip()
    payment_date_str = request.form.get('payment_date')
    payment_date = datetime.strptime(payment_date_str, '%Y-%m-%d').date() if payment_date_str else None
    
    tax_14_enabled = request.form.get('tax_14_enabled') == 'on'
    tax_14_value = float(request.form.get('tax_14_value', 14))
    tax_85_enabled = request.form.get('tax_85_enabled') == 'on'
    tax_85_value = float(request.form.get('tax_85_value', 8.5))
    
    invoice_number = request.form.get('invoice_number', '').strip()
    invoice_date_str = request.form.get('invoice_date')
    invoice_date = datetime.strptime(invoice_date_str, '%Y-%m-%d').date() if invoice_date_str else None
    project_name = request.form.get('project_name', '').strip()
    project_location = request.form.get('project_location', '').strip()
    payment_method = request.form.get('payment_method', '').strip()
    
    # ✅ الحسابات الجديدة
    rental_extra_value = (rental_extra_hours * rental_hour_rate) + (rental_travel_days * travel_rental_rate)
    supply_extra_value = (supply_extra_hours * supply_hour_rate) + (supply_travel_days * travel_supply_rate)
    rental_total = (rental_value * rental_days) + rental_extra_value
    supply_total = (supply_value * supply_days) + supply_extra_value
    
    operation = AlMasaPrivateOperation(
        crane_id=crane_id, start_date=start_date, end_date=end_date,
        supply_value=supply_value, rental_value=rental_value,
        supply_total=supply_total, rental_total=rental_total,
        rental_days=rental_days, supply_days=supply_days,
        rental_extra_hours=rental_extra_hours, rental_hour_rate=rental_hour_rate,
        supply_extra_hours=supply_extra_hours, supply_hour_rate=supply_hour_rate,
        rental_travel_days=rental_travel_days, supply_travel_days=supply_travel_days,
        travel_rental_rate=travel_rental_rate, travel_supply_rate=travel_supply_rate,
        payment_check_number=payment_check_number,
        payment_check_due_date=payment_check_due_date,
        payment_account_name=payment_account_name,
        payment_account_number=payment_account_number,
        payment_date=payment_date,
        tax_14_enabled=tax_14_enabled, tax_14_value=tax_14_value,
        tax_85_enabled=tax_85_enabled, tax_85_value=tax_85_value,
        invoice_number=invoice_number, invoice_date=invoice_date,
        project_name=project_name, project_location=project_location,
        payment_method=payment_method,
        created_by=current_user.id
    )
    db.session.add(operation)
    db.session.commit()
    flash('تم إضافة العملية بنجاح', 'success')
    return redirect(url_for('almasa_private_crane_detail', crane_id=crane_id))


@app.route('/almasa/private-operations/<int:op_id>/delete', methods=['POST'])
@custom_login_required
@role_required('sayed', 'dina', 'admin', 'meg')
def almasa_private_operation_delete(op_id):
    operation = AlMasaPrivateOperation.query.get_or_404(op_id)
    crane_id = operation.crane_id
    AlMasaPrivateCheckOperation.query.filter_by(operation_id=op_id).delete()
    db.session.delete(operation)
    db.session.commit()
    flash('تم حذف العملية بنجاح', 'success')
    return redirect(url_for('almasa_private_crane_detail', crane_id=crane_id))


# ====================================================================
# ✅✅✅ almasa_private_operation_edit — الحسابات الجديدة
# ====================================================================
@app.route('/almasa/private-operations/<int:op_id>/edit', methods=['POST'])
@custom_login_required
@role_required('sayed', 'dina', 'admin', 'meg')
def almasa_private_operation_edit(op_id):
    operation = AlMasaPrivateOperation.query.get_or_404(op_id)
    operation.start_date = datetime.strptime(request.form.get('start_date'), '%Y-%m-%d').date()
    operation.end_date = datetime.strptime(request.form.get('end_date'), '%Y-%m-%d').date() if request.form.get('end_date') else None
    
    operation.supply_value = float(request.form.get('supply_value', 0))
    operation.rental_value = float(request.form.get('rental_value', 0))
    
    operation.rental_days = float(request.form.get('rental_days', 0))
    operation.supply_days = float(request.form.get('supply_days', 0))
    operation.rental_extra_hours = float(request.form.get('rental_extra_hours', 0))
    operation.rental_hour_rate = float(request.form.get('rental_hour_rate', 0))
    operation.supply_extra_hours = float(request.form.get('supply_extra_hours', 0))
    operation.supply_hour_rate = float(request.form.get('supply_hour_rate', 0))
    operation.rental_travel_days = float(request.form.get('rental_travel_days', 0))
    operation.supply_travel_days = float(request.form.get('supply_travel_days', 0))
    operation.travel_rental_rate = float(request.form.get('travel_rental_rate', 0))
    operation.travel_supply_rate = float(request.form.get('travel_supply_rate', 0))
    
    operation.payment_check_number = request.form.get('payment_check_number', '').strip()
    payment_check_due_date_str = request.form.get('payment_check_due_date')
    operation.payment_check_due_date = datetime.strptime(payment_check_due_date_str, '%Y-%m-%d').date() if payment_check_due_date_str else None
    operation.payment_account_name = request.form.get('payment_account_name', '').strip()
    operation.payment_account_number = request.form.get('payment_account_number', '').strip()
    payment_date_str = request.form.get('payment_date')
    operation.payment_date = datetime.strptime(payment_date_str, '%Y-%m-%d').date() if payment_date_str else None
    
    operation.tax_14_enabled = request.form.get('tax_14_enabled') == 'on'
    operation.tax_14_value = float(request.form.get('tax_14_value', 14))
    operation.tax_85_enabled = request.form.get('tax_85_enabled') == 'on'
    operation.tax_85_value = float(request.form.get('tax_85_value', 8.5))
    
    operation.invoice_number = request.form.get('invoice_number', '').strip()
    invoice_date_str = request.form.get('invoice_date')
    operation.invoice_date = datetime.strptime(invoice_date_str, '%Y-%m-%d').date() if invoice_date_str else None
    operation.project_name = request.form.get('project_name', '').strip()
    operation.project_location = request.form.get('project_location', '').strip()
    operation.payment_method = request.form.get('payment_method', '').strip()
    
    # ✅ الحسابات الجديدة
    rental_extra_value = (operation.rental_extra_hours * operation.rental_hour_rate) + (operation.rental_travel_days * operation.travel_rental_rate)
    supply_extra_value = (operation.supply_extra_hours * operation.supply_hour_rate) + (operation.supply_travel_days * operation.travel_supply_rate)
    operation.rental_total = (operation.rental_value * operation.rental_days) + rental_extra_value
    operation.supply_total = (operation.supply_value * operation.supply_days) + supply_extra_value
    
    db.session.commit()
    flash('تم تعديل العملية بنجاح', 'success')
    return redirect(url_for('almasa_private_crane_detail', crane_id=operation.crane_id))


@app.route('/almasa/private-operations/<int:op_id>/update-notes', methods=['POST'])
@custom_login_required
@role_required('sayed', 'dina', 'admin', 'meg')
def almasa_private_operation_update_notes(op_id):
    operation = AlMasaPrivateOperation.query.get_or_404(op_id)
    operation.notes = request.form.get('notes', '')
    db.session.commit()
    flash('تم تحديث الملاحظات بنجاح', 'success')
    return redirect(url_for('almasa_private_crane_detail', crane_id=operation.crane_id))


@app.route('/almasa/private-operations/<int:op_id>/receive-check', methods=['POST'])
@custom_login_required
@role_required('sayed', 'dina', 'admin', 'meg')
def almasa_private_operation_receive_check(op_id):
    operation = AlMasaPrivateOperation.query.get_or_404(op_id)
    check_number = request.form.get('check_number', '').strip()
    issue_date_str = request.form.get('issue_date')
    issue_date = datetime.strptime(issue_date_str, '%Y-%m-%d').date() if issue_date_str else date.today()
    check = AlMasaPrivateCheck(
        crane_id=operation.crane_id, operation_id=op_id,
        check_number=check_number, company_name=operation.project_name,
        amount=operation.rental_total + (operation.rental_total * (operation.tax_14_value / 100) if operation.tax_14_enabled else 0),
        issue_date=issue_date, status='مستلم'
    )
    db.session.add(check)
    operation.check_received = True
    operation.check_received_date = issue_date
    db.session.commit()
    flash('تم تسجيل استلام الشيك بنجاح', 'success')
    return redirect(url_for('almasa_private_crane_detail', crane_id=operation.crane_id))


@app.route('/almasa/private-expenses/add', methods=['POST'])
@custom_login_required
@role_required('sayed', 'dina', 'admin', 'meg')
def almasa_add_private_expense():
    crane_id = int(request.form.get('crane_id'))
    operation_id = request.form.get('operation_id')
    operation_id = int(operation_id) if operation_id else None
    date_str = datetime.strptime(request.form.get('date'), '%Y-%m-%d').date()
    expense_type = request.form.get('expense_type')
    amount = float(request.form.get('amount', 0))
    paid_by = request.form.get('paid_by', '').strip()
    notes = request.form.get('notes')
    expense = AlMasaPrivateExpense(crane_id=crane_id, operation_id=operation_id, date=date_str,
                                   expense_type=expense_type, amount=amount, paid_by=paid_by,
                                   notes=notes, created_by=current_user.id)
    db.session.add(expense)
    db.session.commit()
    flash('تم إضافة المصروف بنجاح', 'success')
    return redirect(url_for('almasa_private_crane_detail', crane_id=crane_id))


@app.route('/almasa/private-expenses/<int:exp_id>/delete', methods=['POST'])
@custom_login_required
@role_required('sayed', 'dina', 'admin', 'meg')
def almasa_private_expense_delete(exp_id):
    expense = AlMasaPrivateExpense.query.get_or_404(exp_id)
    crane_id = expense.crane_id
    db.session.delete(expense)
    db.session.commit()
    flash('تم حذف المصروف بنجاح', 'success')
    return redirect(url_for('almasa_private_crane_detail', crane_id=crane_id))


@app.route('/almasa/private-expenses/<int:exp_id>/edit', methods=['POST'])
@custom_login_required
@role_required('sayed', 'dina', 'admin', 'meg')
def almasa_private_expense_edit(exp_id):
    expense = AlMasaPrivateExpense.query.get_or_404(exp_id)
    expense.date = datetime.strptime(request.form.get('date'), '%Y-%m-%d').date()
    expense.expense_type = request.form.get('expense_type')
    expense.amount = float(request.form.get('amount', 0))
    expense.paid_by = request.form.get('paid_by', '').strip()
    expense.notes = request.form.get('notes')
    db.session.commit()
    flash('تم تعديل المصروف بنجاح', 'success')
    return redirect(url_for('almasa_private_crane_detail', crane_id=expense.crane_id))


@app.route('/almasa/private-checks/<int:check_id>/delete', methods=['POST'])
@custom_login_required
@role_required('sayed', 'dina', 'admin', 'meg')
def almasa_private_check_delete(check_id):
    check = AlMasaPrivateCheck.query.get_or_404(check_id)
    crane_id = check.crane_id
    AlMasaPrivateCheckOperation.query.filter_by(check_id=check_id).delete()
    db.session.delete(check)
    db.session.commit()
    flash('تم حذف الشيك بنجاح', 'success')
    return redirect(url_for('almasa_private_crane_detail', crane_id=crane_id))


@app.route('/almasa/private-checks/<int:check_id>/report')
@custom_login_required
@role_required('sayed', 'dina', 'admin', 'meg')
def almasa_private_check_report(check_id):
    check = AlMasaPrivateCheck.query.get_or_404(check_id)
    crane = AlMasaPrivateCrane.query.get(check.crane_id)
    check_operations = AlMasaPrivateCheckOperation.query.filter_by(check_id=check_id).all()
    operations_list = []
    total_rental = 0
    total_supply = 0
    total_days = 0
    for co in check_operations:
        operation = AlMasaPrivateOperation.query.get(co.operation_id)
        if operation:
            operations_list.append(operation)
            total_rental += operation.rental_total
            total_supply += operation.supply_total
            total_days += operation.days_count
    if operations_list:
        min_date = min(o.start_date for o in operations_list)
        dates_with_end = [o.end_date for o in operations_list if o.end_date]
        max_date = max(dates_with_end) if dates_with_end else max(o.start_date for o in operations_list)
        expenses = AlMasaPrivateExpense.query.filter(
            AlMasaPrivateExpense.crane_id == check.crane_id,
            AlMasaPrivateExpense.date >= min_date,
            AlMasaPrivateExpense.date <= max_date
        ).all()
    else:
        expenses = []
    total_expenses = sum(e.amount for e in expenses)
    total_supply = total_rental
    check_amount = total_supply + (total_supply * (tax_14_value / 100))
    tax_14 = total_supply * (tax_14_value / 100)
    admin_profit = total_rental - total_supply
    supply_profit = total_supply - total_expenses
    owner_share = admin_profit + supply_profit
    return render_template('almasa/check_report.html',
                           check=check, crane=crane,
                           operations=operations_list, expenses=expenses,
                           total_rental=total_rental, total_supply=total_supply,
                           total_days=total_days, total_expenses=total_expenses,
                           check_amount=check_amount, tax_14=tax_14,
                           admin_profit=admin_profit, supply_profit=supply_profit,
                           owner_share=owner_share, report_type='private')


# ====================================================================
# ==================== النوع 3: توريدات ====================
# ====================================================================
@app.route('/almasa/supplies', methods=['GET', 'POST'])
@custom_login_required
@role_required('sayed', 'dina', 'admin', 'meg')
def almasa_supplies():
    if request.method == 'POST':
        name = request.form.get('name')
        supplier_name = request.form.get('supplier_name')
        customer_name = request.form.get('customer_name')
        notes = request.form.get('notes')
        new_supply = AlMasaSupply(name=name, supplier_name=supplier_name,
                                   customer_name=customer_name, notes=notes)
        db.session.add(new_supply)
        db.session.commit()
        flash('تم إضافة التوريدة بنجاح', 'success')
        return redirect(url_for('almasa_supplies'))
    supplies = AlMasaSupply.query.order_by(AlMasaSupply.id.asc()).all()
    return render_template('almasa/supplies.html', supplies=supplies)


@app.route('/almasa/supplies/<int:supply_id>')
@custom_login_required
@role_required('sayed', 'dina', 'admin', 'meg')
def almasa_supply_detail(supply_id):
    supply = AlMasaSupply.query.get_or_404(supply_id)
    operations = AlMasaSupplyOperation.query.filter_by(supply_id=supply_id).order_by(AlMasaSupplyOperation.start_date.asc()).all()
    expenses = AlMasaSupplyExpense.query.filter_by(supply_id=supply_id).order_by(AlMasaSupplyExpense.date.asc()).all()
    total_rental = sum(o.rental_total or 0 for o in operations)
    total_supply = sum(o.supply_total or 0 for o in operations)
    total_expenses = sum(e.amount or 0 for e in expenses)
    total_paid = sum(o.rental_total or 0 for o in operations if o.check_received)
    total_unpaid = sum(o.rental_total or 0 for o in operations if not o.check_received)
    total_operations = total_rental
    return render_template('almasa/supply_detail.html',
                           supply=supply, operations=operations, expenses=expenses,
                           total_rental=total_rental, total_supply=total_supply,
                           total_expenses=total_expenses,
                           total_paid=total_paid, total_unpaid=total_unpaid,
                           total_operations=total_operations)


@app.route('/almasa/supplies/<int:supply_id>/delete', methods=['POST'])
@custom_login_required
@role_required('sayed', 'dina', 'admin', 'meg')
def almasa_supply_delete(supply_id):
    supply = AlMasaSupply.query.get_or_404(supply_id)
    db.session.delete(supply)
    db.session.commit()
    flash('تم حذف التوريدة بنجاح', 'success')
    return redirect(url_for('almasa_supplies'))


# ====================================================================
# ✅✅✅ almasa_add_supply_operation — الحسابات الجديدة
# ====================================================================
@app.route('/almasa/supply-operations/add', methods=['POST'])
@custom_login_required
@role_required('sayed', 'dina', 'admin', 'meg')
def almasa_add_supply_operation():
    supply_id = int(request.form.get('supply_id'))
    start_date = datetime.strptime(request.form.get('start_date'), '%Y-%m-%d').date()
    end_date = datetime.strptime(request.form.get('end_date'), '%Y-%m-%d').date() if request.form.get('end_date') else None
    
    supply_value = float(request.form.get('supply_value', 0))
    rental_value = float(request.form.get('rental_value', 0))
    
    rental_days = float(request.form.get('rental_days', 0))
    supply_days = float(request.form.get('supply_days', 0))
    rental_extra_hours = float(request.form.get('rental_extra_hours', 0))
    rental_hour_rate = float(request.form.get('rental_hour_rate', 0))
    supply_extra_hours = float(request.form.get('supply_extra_hours', 0))
    supply_hour_rate = float(request.form.get('supply_hour_rate', 0))
    rental_travel_days = float(request.form.get('rental_travel_days', 0))
    supply_travel_days = float(request.form.get('supply_travel_days', 0))
    travel_rental_rate = float(request.form.get('travel_rental_rate', 0))
    travel_supply_rate = float(request.form.get('travel_supply_rate', 0))
    
    payment_check_number = request.form.get('payment_check_number', '').strip()
    payment_check_due_date_str = request.form.get('payment_check_due_date')
    payment_check_due_date = datetime.strptime(payment_check_due_date_str, '%Y-%m-%d').date() if payment_check_due_date_str else None
    payment_account_name = request.form.get('payment_account_name', '').strip()
    payment_account_number = request.form.get('payment_account_number', '').strip()
    payment_date_str = request.form.get('payment_date')
    payment_date = datetime.strptime(payment_date_str, '%Y-%m-%d').date() if payment_date_str else None
    
    tax_14_enabled = request.form.get('tax_14_enabled') == 'on'
    tax_14_value = float(request.form.get('tax_14_value', 14))
    tax_85_enabled = request.form.get('tax_85_enabled') == 'on'
    tax_85_value = float(request.form.get('tax_85_value', 8.5))
    
    invoice_number = request.form.get('invoice_number', '').strip()
    invoice_date_str = request.form.get('invoice_date')
    invoice_date = datetime.strptime(invoice_date_str, '%Y-%m-%d').date() if invoice_date_str else None
    project_name = request.form.get('project_name', '').strip()
    project_location = request.form.get('project_location', '').strip()
    payment_method = request.form.get('payment_method', '').strip()
    
    # ✅ الحسابات الجديدة
    rental_extra_value = (rental_extra_hours * rental_hour_rate) + (rental_travel_days * travel_rental_rate)
    supply_extra_value = (supply_extra_hours * supply_hour_rate) + (supply_travel_days * travel_supply_rate)
    rental_total = (rental_value * rental_days) + rental_extra_value
    supply_total = (supply_value * supply_days) + supply_extra_value
    
    operation = AlMasaSupplyOperation(
        supply_id=supply_id, start_date=start_date, end_date=end_date,
        supply_value=supply_value, rental_value=rental_value,
        supply_total=supply_total, rental_total=rental_total,
        rental_days=rental_days, supply_days=supply_days,
        rental_extra_hours=rental_extra_hours, rental_hour_rate=rental_hour_rate,
        supply_extra_hours=supply_extra_hours, supply_hour_rate=supply_hour_rate,
        rental_travel_days=rental_travel_days, supply_travel_days=supply_travel_days,
        travel_rental_rate=travel_rental_rate, travel_supply_rate=travel_supply_rate,
        payment_check_number=payment_check_number,
        payment_check_due_date=payment_check_due_date,
        payment_account_name=payment_account_name,
        payment_account_number=payment_account_number,
        payment_date=payment_date,
        tax_14_enabled=tax_14_enabled, tax_14_value=tax_14_value,
        tax_85_enabled=tax_85_enabled, tax_85_value=tax_85_value,
        invoice_number=invoice_number, invoice_date=invoice_date,
        project_name=project_name, project_location=project_location,
        payment_method=payment_method,
        created_by=current_user.id
    )
    db.session.add(operation)
    db.session.commit()
    flash('تم إضافة العملية بنجاح', 'success')
    return redirect(url_for('almasa_supply_detail', supply_id=supply_id))


@app.route('/almasa/supply-operations/<int:op_id>/delete', methods=['POST'])
@custom_login_required
@role_required('sayed', 'dina', 'admin', 'meg')
def almasa_supply_operation_delete(op_id):
    operation = AlMasaSupplyOperation.query.get_or_404(op_id)
    supply_id = operation.supply_id
    db.session.delete(operation)
    db.session.commit()
    flash('تم حذف العملية بنجاح', 'success')
    return redirect(url_for('almasa_supply_detail', supply_id=supply_id))


# ====================================================================
# ✅✅✅ almasa_supply_operation_edit — الحسابات الجديدة
# ====================================================================
@app.route('/almasa/supply-operations/<int:op_id>/edit', methods=['POST'])
@custom_login_required
@role_required('sayed', 'dina', 'admin', 'meg')
def almasa_supply_operation_edit(op_id):
    operation = AlMasaSupplyOperation.query.get_or_404(op_id)
    operation.start_date = datetime.strptime(request.form.get('start_date'), '%Y-%m-%d').date()
    operation.end_date = datetime.strptime(request.form.get('end_date'), '%Y-%m-%d').date() if request.form.get('end_date') else None
    
    operation.supply_value = float(request.form.get('supply_value', 0))
    operation.rental_value = float(request.form.get('rental_value', 0))
    
    operation.rental_days = float(request.form.get('rental_days', 0))
    operation.supply_days = float(request.form.get('supply_days', 0))
    operation.rental_extra_hours = float(request.form.get('rental_extra_hours', 0))
    operation.rental_hour_rate = float(request.form.get('rental_hour_rate', 0))
    operation.supply_extra_hours = float(request.form.get('supply_extra_hours', 0))
    operation.supply_hour_rate = float(request.form.get('supply_hour_rate', 0))
    operation.rental_travel_days = float(request.form.get('rental_travel_days', 0))
    operation.supply_travel_days = float(request.form.get('supply_travel_days', 0))
    operation.travel_rental_rate = float(request.form.get('travel_rental_rate', 0))
    operation.travel_supply_rate = float(request.form.get('travel_supply_rate', 0))
    
    operation.payment_check_number = request.form.get('payment_check_number', '').strip()
    payment_check_due_date_str = request.form.get('payment_check_due_date')
    operation.payment_check_due_date = datetime.strptime(payment_check_due_date_str, '%Y-%m-%d').date() if payment_check_due_date_str else None
    operation.payment_account_name = request.form.get('payment_account_name', '').strip()
    operation.payment_account_number = request.form.get('payment_account_number', '').strip()
    payment_date_str = request.form.get('payment_date')
    operation.payment_date = datetime.strptime(payment_date_str, '%Y-%m-%d').date() if payment_date_str else None
    
    operation.tax_14_enabled = request.form.get('tax_14_enabled') == 'on'
    operation.tax_14_value = float(request.form.get('tax_14_value', 14))
    operation.tax_85_enabled = request.form.get('tax_85_enabled') == 'on'
    operation.tax_85_value = float(request.form.get('tax_85_value', 8.5))
    
    operation.invoice_number = request.form.get('invoice_number', '').strip()
    invoice_date_str = request.form.get('invoice_date')
    operation.invoice_date = datetime.strptime(invoice_date_str, '%Y-%m-%d').date() if invoice_date_str else None
    operation.project_name = request.form.get('project_name', '').strip()
    operation.project_location = request.form.get('project_location', '').strip()
    operation.payment_method = request.form.get('payment_method', '').strip()
    
    # ✅ الحسابات الجديدة
    rental_extra_value = (operation.rental_extra_hours * operation.rental_hour_rate) + (operation.rental_travel_days * operation.travel_rental_rate)
    supply_extra_value = (operation.supply_extra_hours * operation.supply_hour_rate) + (operation.supply_travel_days * operation.travel_supply_rate)
    operation.rental_total = (operation.rental_value * operation.rental_days) + rental_extra_value
    operation.supply_total = (operation.supply_value * operation.supply_days) + supply_extra_value
    
    db.session.commit()
    flash('تم تعديل العملية بنجاح', 'success')
    return redirect(url_for('almasa_supply_detail', supply_id=operation.supply_id))


@app.route('/almasa/supply-operations/<int:op_id>/update-notes', methods=['POST'])
@custom_login_required
@role_required('sayed', 'dina', 'admin', 'meg')
def almasa_supply_operation_update_notes(op_id):
    operation = AlMasaSupplyOperation.query.get_or_404(op_id)
    operation.notes = request.form.get('notes', '')
    db.session.commit()
    flash('تم تحديث الملاحظات بنجاح', 'success')
    return redirect(url_for('almasa_supply_detail', supply_id=operation.supply_id))


@app.route('/almasa/supply-operations/<int:op_id>/receive-check', methods=['POST'])
@custom_login_required
@role_required('sayed', 'dina', 'admin', 'meg')
def almasa_supply_operation_receive_check(op_id):
    operation = AlMasaSupplyOperation.query.get_or_404(op_id)
    issue_date_str = request.form.get('issue_date')
    issue_date = datetime.strptime(issue_date_str, '%Y-%m-%d').date() if issue_date_str else date.today()
    operation.check_received = True
    operation.check_received_date = issue_date
    db.session.commit()
    flash('تم تسجيل استلام الشيك بنجاح', 'success')
    return redirect(url_for('almasa_supply_detail', supply_id=operation.supply_id))


@app.route('/almasa/supply-expenses/add', methods=['POST'])
@custom_login_required
@role_required('sayed', 'dina', 'admin', 'meg')
def almasa_add_supply_expense():
    supply_id = int(request.form.get('supply_id'))
    operation_id = request.form.get('operation_id')
    operation_id = int(operation_id) if operation_id else None
    date_str = datetime.strptime(request.form.get('date'), '%Y-%m-%d').date()
    expense_type = request.form.get('expense_type')
    amount = float(request.form.get('amount', 0))
    paid_by = request.form.get('paid_by', '').strip()
    notes = request.form.get('notes')
    expense = AlMasaSupplyExpense(supply_id=supply_id, operation_id=operation_id, date=date_str,
                                  expense_type=expense_type, amount=amount, paid_by=paid_by,
                                  notes=notes, created_by=current_user.id)
    db.session.add(expense)
    db.session.commit()
    flash('تم إضافة المصروف بنجاح', 'success')
    return redirect(url_for('almasa_supply_detail', supply_id=supply_id))


@app.route('/almasa/supply-expenses/<int:exp_id>/delete', methods=['POST'])
@custom_login_required
@role_required('sayed', 'dina', 'admin', 'meg')
def almasa_supply_expense_delete(exp_id):
    expense = AlMasaSupplyExpense.query.get_or_404(exp_id)
    supply_id = expense.supply_id
    db.session.delete(expense)
    db.session.commit()
    flash('تم حذف المصروف بنجاح', 'success')
    return redirect(url_for('almasa_supply_detail', supply_id=supply_id))


@app.route('/almasa/supply-expenses/<int:exp_id>/edit', methods=['POST'])
@custom_login_required
@role_required('sayed', 'dina', 'admin', 'meg')
def almasa_supply_expense_edit(exp_id):
    expense = AlMasaSupplyExpense.query.get_or_404(exp_id)
    expense.date = datetime.strptime(request.form.get('date'), '%Y-%m-%d').date()
    expense.expense_type = request.form.get('expense_type')
    expense.amount = float(request.form.get('amount', 0))
    expense.paid_by = request.form.get('paid_by', '').strip()
    expense.notes = request.form.get('notes')
    db.session.commit()
    flash('تم تعديل المصروف بنجاح', 'success')
    return redirect(url_for('almasa_supply_detail', supply_id=expense.supply_id))


@app.route('/almasa/supplies/<int:supply_id>/report')
@custom_login_required
@role_required('sayed', 'dina', 'admin', 'meg')
def almasa_supply_report(supply_id):
    supply = AlMasaSupply.query.get_or_404(supply_id)
    operations = AlMasaSupplyOperation.query.filter_by(supply_id=supply_id).order_by(AlMasaSupplyOperation.start_date.asc()).all()
    expenses = AlMasaSupplyExpense.query.filter_by(supply_id=supply_id).order_by(AlMasaSupplyExpense.date.asc()).all()
    total_rental = sum(o.rental_total or 0 for o in operations)
    total_supply = sum(o.supply_total or 0 for o in operations)
    total_expenses = sum(e.amount or 0 for e in expenses)
    first_op = operations[0] if operations else None
    tax_85_value = first_op.tax_85_value if first_op and first_op.tax_85_enabled else 0
    tax_85_amount = total_rental * (tax_85_value / 100)
    admin_profit = total_supply - total_rental - total_expenses - tax_85_amount
    check_amount = total_suuply
    tax_14 = total_rental * ((first_op.tax_14_value / 100) if first_op and first_op.tax_14_enabled else 0)
    return render_template('almasa/supply_report.html',
                           supply=supply, operations=operations, expenses=expenses,
                           total_rental=total_rental, total_supply=total_supply,
                           total_expenses=total_expenses, check_amount=check_amount,
                           tax_14=tax_14, tax_85_amount=tax_85_amount,
                           admin_profit=admin_profit)
# ====================================================================
# ==================== تقارير الماسة (لكل ونش) ====================
# ====================================================================
@app.route('/almasa/reports/crane-profit/<int:crane_id>')
@custom_login_required
@role_required('sayed', 'dina', 'admin', 'meg')
def almasa_crane_profit_report(crane_id):
    crane = AlMasaCrane.query.get_or_404(crane_id)
    operations = AlMasaOperation.query.filter_by(crane_id=crane_id).all()
    expenses = AlMasaExpense.query.filter_by(crane_id=crane_id).all()
    partners = AlMasaPartner.query.filter_by(crane_id=crane_id).all()
    total_rental = sum(o.rental_total or 0 for o in operations)
    total_supply = sum(o.supply_total or 0 for o in operations)
    total_expenses = sum(e.amount or 0 for e in expenses)
    total_days = sum(o.days_count or 0 for o in operations)
    tax_14_amount = sum((o.rental_total * o.tax_14_value / 100) for o in operations if o.tax_14_enabled)
    tax_85_amount = sum((o.supply_total * o.tax_85_value / 100) for o in operations if o.tax_85_enabled)
    admin_profit = total_rental - total_supply
    supply_profit = total_supply - total_expenses - tax_85_amount
    total_profit = admin_profit + supply_profit
    basic_partners = [p for p in partners if p.is_basic]
    basic_count = len(basic_partners) if basic_partners else 1
    basic_diff = admin_profit / basic_count
    partners_profit = []
    total_partners_share = 0
    for p in partners:
        if p.is_basic:
            actual_share = basic_diff
            default_share = supply_profit * (p.percentage / 100)
        else:
            actual_share = 0
            default_share = supply_profit * (p.percentage / 100)
        total_share = actual_share + default_share
        total_partners_share += total_share
        partners_profit.append({
            'partner': p, 'actual_share': actual_share,
            'default_share': default_share, 'total_share': total_share
        })
    return render_template('almasa/crane_profit_report.html',
                           crane=crane, operations=operations,
                           partners=partners, expenses=expenses,
                           total_days=total_days,
                           total_rental=total_rental, total_supply=total_supply,
                           total_expenses=total_expenses,
                           tax_14_amount=tax_14_amount, tax_85_amount=tax_85_amount,
                           admin_profit=admin_profit, supply_profit=supply_profit,
                           total_profit=total_profit,
                           partners_profit=partners_profit,
                           total_partners_share=total_partners_share)


@app.route('/almasa/reports/private-crane-profit/<int:crane_id>')
@custom_login_required
@role_required('sayed', 'dina', 'admin', 'meg')
def almasa_private_crane_profit_report(crane_id):
    crane = AlMasaPrivateCrane.query.get_or_404(crane_id)
    operations = AlMasaPrivateOperation.query.filter_by(crane_id=crane_id).all()
    expenses = AlMasaPrivateExpense.query.filter_by(crane_id=crane_id).all()
    total_rental = sum(o.rental_total or 0 for o in operations)
    total_expenses = sum(e.amount or 0 for e in expenses)
    total_days = sum(o.days_count or 0 for o in operations)
    tax_14_amount = sum((o.rental_total * o.tax_14_value / 100) for o in operations if o.tax_14_enabled)
    tax_85_amount = sum((o.rental_total * o.tax_85_value / 100) for o in operations if o.tax_85_enabled)
    owner_share = total_rental - total_expenses - tax_85_amount
    return render_template('almasa/private_crane_profit_report.html',
                           crane=crane, operations=operations,
                           total_days=total_days,
                           total_rental=total_rental,
                           total_expenses=total_expenses,
                           tax_14_amount=tax_14_amount,
                           tax_85_amount=tax_85_amount,
                           owner_share=owner_share)


@app.route('/almasa/reports/supply-profit/<int:supply_id>')
@custom_login_required
@role_required('sayed', 'dina', 'admin', 'meg')
def almasa_supply_profit_single(supply_id):
    supply = AlMasaSupply.query.get_or_404(supply_id)
    operations = AlMasaSupplyOperation.query.filter_by(supply_id=supply_id).all()
    expenses = AlMasaSupplyExpense.query.filter_by(supply_id=supply_id).all()
    total_rental = sum(o.rental_total or 0 for o in operations)
    total_supply = sum(o.supply_total or 0 for o in operations)
    total_expenses = sum(e.amount or 0 for e in expenses)
    tax_14_amount = sum((o.rental_total * o.tax_14_value / 100) for o in operations if o.tax_14_enabled)
    tax_85_amount = sum((o.rental_total * o.tax_85_value / 100) for o in operations if o.tax_85_enabled)
    admin_profit = total_rental - total_supply - total_expenses - tax_85_amount
    return render_template('almasa/supply_profit_single.html',
                           supply=supply, operations=operations,
                           total_rental=total_rental, total_supply=total_supply,
                           total_expenses=total_expenses,
                           tax_14_amount=tax_14_amount, tax_85_amount=tax_85_amount,
                           admin_profit=admin_profit)


# ====================================================================
# ==================== تقارير الماسة (الكل) ====================
# ====================================================================
@app.route('/almasa/reports/admin-profit')
@custom_login_required
@role_required('sayed', 'dina', 'admin', 'meg')
def almasa_admin_profit_report():
    partnership_cranes = AlMasaCrane.query.filter_by(crane_type='partnership').all()
    private_cranes = AlMasaPrivateCrane.query.all()
    supplies = AlMasaSupply.query.all()
    partnership_details = []
    partnership_admin_profit = 0
    for crane in partnership_cranes:
        total_rental = sum(o.rental_total or 0 for o in crane.operations)
        total_supply = sum(o.supply_total or 0 for o in crane.operations)
        admin_profit = total_rental - total_supply
        partnership_admin_profit += admin_profit
        partnership_details.append({
            'crane': crane, 'total_rental': total_rental,
            'total_supply': total_supply, 'admin_profit': admin_profit
        })
    private_details = []
    private_admin_profit = 0
    for crane in private_cranes:
        total_rental = sum(o.rental_total or 0 for o in crane.operations)
        total_supply = sum(o.supply_total or 0 for o in crane.operations)
        admin_profit = total_rental - total_supply
        private_admin_profit += admin_profit
        private_details.append({
            'crane': crane, 'total_rental': total_rental,
            'admin_profit': admin_profit
        })
    supply_details = []
    supply_admin_profit = 0
    for supply in supplies:
        total_rental = sum(o.rental_total or 0 for o in supply.operations)
        total_supply = sum(o.supply_total or 0 for o in supply.operations)
        total_expenses = sum(e.amount or 0 for e in supply.expenses)
        tax_85_amount = sum((o.rental_total * o.tax_85_value / 100) for o in supply.operations if o.tax_85_enabled)
        admin_profit = total_rental - total_supply - total_expenses - tax_85_amount
        supply_admin_profit += admin_profit
        supply_details.append({
            'supply': supply, 'total_rental': total_rental,
            'total_supply': total_supply, 'admin_profit': admin_profit
        })
    grand_admin_profit = partnership_admin_profit + private_admin_profit + supply_admin_profit
    return render_template('almasa/admin_profit_report.html',
                           partnership_cranes=partnership_cranes,
                           private_cranes=private_cranes,
                           supplies=supplies,
                           partnership_details=partnership_details,
                           private_details=private_details,
                           supply_details=supply_details,
                           partnership_admin_profit=partnership_admin_profit,
                           private_admin_profit=private_admin_profit,
                           supply_admin_profit=supply_admin_profit,
                           grand_admin_profit=grand_admin_profit)


@app.route('/almasa/reports/supply-profit-all')
@custom_login_required
@role_required('sayed', 'dina', 'admin', 'meg')
def almasa_supply_profit_report():
    partnership_cranes = AlMasaCrane.query.filter_by(crane_type='partnership').all()
    private_cranes = AlMasaPrivateCrane.query.all()
    partnership_details = []
    partnership_supply_profit = 0
    for crane in partnership_cranes:
        total_supply = sum(o.supply_total or 0 for o in crane.operations)
        total_expenses = sum(e.amount or 0 for e in crane.expenses)
        tax_85_amount = sum((o.supply_total * o.tax_85_value / 100) for o in crane.operations if o.tax_85_enabled)
        supply_profit = total_supply - total_expenses - tax_85_amount
        partnership_supply_profit += supply_profit
        partnership_details.append({
            'crane': crane, 'total_supply': total_supply,
            'total_expenses': total_expenses,
            'tax_85_amount': tax_85_amount, 'supply_profit': supply_profit
        })
    private_details = []
    private_supply_profit = 0
    for crane in private_cranes:
        total_supply = sum(o.supply_total or 0 for o in crane.operations)
        total_expenses = sum(e.amount or 0 for e in crane.expenses)
        tax_85_amount = sum((o.supply_total * o.tax_85_value / 100) for o in crane.operations if o.tax_85_enabled)
        supply_profit = total_supply - total_expenses - tax_85_amount
        private_supply_profit += supply_profit
        private_details.append({
            'crane': crane, 'total_supply': total_supply,
            'total_expenses': total_expenses,
            'tax_85_amount': tax_85_amount, 'supply_profit': supply_profit
        })
    grand_supply_profit = partnership_supply_profit + private_supply_profit
    return render_template('almasa/supply_profit_report.html',
                           partnership_cranes=partnership_cranes,
                           private_cranes=private_cranes,
                           partnership_details=partnership_details,
                           private_details=private_details,
                           partnership_supply_profit=partnership_supply_profit,
                           private_supply_profit=private_supply_profit,
                           grand_supply_profit=grand_supply_profit)


# ====================================================================
# ==================== تقارير الماسة (العامة) ====================
# ====================================================================
@app.route('/almasa/reports')
@custom_login_required
@role_required('sayed', 'dina', 'admin', 'meg')
def almasa_reports():
    partnership_cranes = AlMasaCrane.query.filter_by(crane_type='partnership').all()
    private_cranes = AlMasaPrivateCrane.query.all()
    supplies = AlMasaSupply.query.all()
    return render_template('almasa/reports.html',
                           partnership_cranes=partnership_cranes,
                           private_cranes=private_cranes, supplies=supplies)


@app.route('/almasa/reports/crane/<int:crane_id>')
@custom_login_required
@role_required('sayed', 'dina', 'admin', 'meg')
def almasa_crane_report(crane_id):
    crane = AlMasaCrane.query.get_or_404(crane_id)
    operations = AlMasaOperation.query.filter_by(crane_id=crane_id).order_by(AlMasaOperation.start_date.asc()).all()
    expenses = AlMasaExpense.query.filter_by(crane_id=crane_id).order_by(AlMasaExpense.date.asc()).all()
    partners = AlMasaPartner.query.filter_by(crane_id=crane_id).all()
    total_rental = sum(o.rental_total or 0 for o in operations)
    total_supply = sum(o.supply_total or 0 for o in operations)
    total_expenses = sum(e.amount or 0 for e in expenses)
    non_basic = [p for p in partners if not p.is_basic]
    if len(non_basic) == 0:
        total_supply = total_rental
    admin_profit = total_rental - total_supply
    supply_profit = total_supply - total_expenses
    basic_partners = [p for p in partners if p.is_basic]
    basic_count = len(basic_partners) if basic_partners else 1
    basic_diff = admin_profit / basic_count
    partners_profit = []
    for p in partners:
        if p.is_basic:
            actual_share = basic_diff
            default_share = supply_profit * (p.percentage / 100)
        else:
            actual_share = 0
            default_share = supply_profit * (p.percentage / 100)
        partners_profit.append({
            'partner': p, 'actual_share': actual_share,
            'default_share': default_share, 'total_share': actual_share + default_share
        })
    return render_template('almasa/crane_report.html',
                           crane=crane, operations=operations, expenses=expenses,
                           partners=partners, total_rental=total_rental,
                           total_supply=total_supply, total_expenses=total_expenses,
                           admin_profit=admin_profit, supply_profit=supply_profit,
                           basic_diff=basic_diff, partners_profit=partners_profit,
                           report_type='partnership')


@app.route('/almasa/reports/private-crane/<int:crane_id>')
@custom_login_required
@role_required('sayed', 'dina', 'admin', 'meg')
def almasa_private_crane_report(crane_id):
    crane = AlMasaPrivateCrane.query.get_or_404(crane_id)
    operations = AlMasaPrivateOperation.query.filter_by(crane_id=crane_id).order_by(AlMasaPrivateOperation.start_date.asc()).all()
    expenses = AlMasaPrivateExpense.query.filter_by(crane_id=crane_id).order_by(AlMasaPrivateExpense.date.asc()).all()
    total_rental = sum(o.rental_total or 0 for o in operations)
    total_supply = total_rental
    total_expenses = sum(e.amount or 0 for e in expenses)
    admin_profit = total_rental - total_supply
    supply_profit = total_supply - total_expenses
    owner_share = admin_profit + supply_profit
    return render_template('almasa/private_crane_report.html',
                           crane=crane, operations=operations, expenses=expenses,
                           total_rental=total_rental, total_supply=total_supply,
                           total_expenses=total_expenses,
                           admin_profit=admin_profit, supply_profit=supply_profit,
                           owner_share=owner_share)


# ====================================================================
# ==================== حساب الشريك / العميل ====================
# ====================================================================
@app.route('/almasa/partner-accounts', methods=['GET', 'POST'])
@custom_login_required
@role_required('sayed', 'dina', 'admin', 'meg')
def almasa_partner_accounts():
    if request.method == 'POST':
        if request.form.get('delete_id'):
            record_id = int(request.form.get('delete_id'))
            record = AlMasaPartnerAccount.query.get_or_404(record_id)
            db.session.delete(record)
            db.session.commit()
            flash('تم الحذف بنجاح', 'success')
            return redirect(url_for('almasa_partner_accounts'))
        if request.form.get('edit_id'):
            record_id = int(request.form.get('edit_id'))
            record = AlMasaPartnerAccount.query.get_or_404(record_id)
            record.person_name = request.form.get('person_name')
            record.person_type = request.form.get('person_type')
            record.date = datetime.strptime(request.form.get('date'), '%Y-%m-%d').date()
            record.amount = float(request.form.get('amount', 0))
            record.description = request.form.get('description')
            record.notes = request.form.get('notes')
            db.session.commit()
            flash('تم التعديل بنجاح', 'success')
            return redirect(url_for('almasa_partner_accounts'))
        person_name = request.form.get('person_name')
        person_type = request.form.get('person_type', 'شريك')
        date_str = datetime.strptime(request.form.get('date'), '%Y-%m-%d').date()
        amount = float(request.form.get('amount', 0))
        description = request.form.get('description')
        notes = request.form.get('notes')
        new_record = AlMasaPartnerAccount(person_name=person_name, person_type=person_type,
                                          date=date_str, amount=amount, description=description,
                                          notes=notes, created_by=current_user.id)
        db.session.add(new_record)
        db.session.commit()
        flash('تم إضافة السجل بنجاح', 'success')
        return redirect(url_for('almasa_partner_accounts'))
    records = AlMasaPartnerAccount.query.order_by(AlMasaPartnerAccount.date.desc()).all()
    return render_template('almasa/partner_accounts.html', records=records)


@app.route('/almasa/partner-accounts/person/<string:person_name>')
@custom_login_required
@role_required('sayed', 'dina', 'admin', 'meg')
def almasa_person_report(person_name):
    records = AlMasaPartnerAccount.query.filter_by(person_name=person_name).order_by(AlMasaPartnerAccount.date.asc()).all()
    total = sum(r.amount for r in records)
    return render_template('almasa/person_report.html',
                           person_name=person_name, records=records, total=total)


# ====================================================================
# ==================== العمليات غير المدفوعة ====================
# ====================================================================
@app.route('/almasa/unpaid-operations')
@custom_login_required
@role_required('sayed', 'dina', 'admin', 'meg')
def almasa_unpaid_operations():
    partnership_ops = AlMasaOperation.query.filter(
        (AlMasaOperation.check_received == False) | (AlMasaOperation.check_received == None)
    ).all()
    private_ops = AlMasaPrivateOperation.query.filter(
        (AlMasaPrivateOperation.check_received == False) | (AlMasaPrivateOperation.check_received == None)
    ).all()
    supply_ops = AlMasaSupplyOperation.query.filter(
        (AlMasaSupplyOperation.check_received == False) | (AlMasaSupplyOperation.check_received == None)
    ).all()
    partnership_total = sum(o.rental_total or 0 for o in partnership_ops)
    private_total = sum(o.rental_total or 0 for o in private_ops)
    supply_total = sum(o.rental_total or 0 for o in supply_ops)
    grand_total = partnership_total + private_total + supply_total
    return render_template('almasa/unpaid_operations.html',
                           partnership_ops=partnership_ops,
                           private_ops=private_ops,
                           supply_ops=supply_ops,
                           partnership_total=partnership_total,
                           private_total=private_total,
                           supply_total=supply_total,
                           grand_total=grand_total)


# ==================== التشغيل ====================
if __name__ == '__main__':
    app.run(host='0.0.0.0', port=5000, debug=True)    
    
